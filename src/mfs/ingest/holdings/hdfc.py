"""HDFC monthly portfolio holdings adapter (Phase 3.C).

HDFC publishes one Excel per scheme per month at:
    https://files.hdfcfund.com/s3fs-public/{YYYY-MM}/Monthly%20{SchemeName}
        %20-%20{DD}%20{Month}%20{YYYY}.xlsx

URL convention: publish month = data month + 1 (April 2026 data → 2026-05
folder, same as ABSL/Mirae/SBI). The disclosures index page itself is a
Next.js SPA but the rendered HTML carries every Excel link as a plain
absolute URL — so a single page scrape gives us the full month's catalog.

Excel structure (validated against HDFC Flexi Cap, April 2026):
- Sheet 0: equity holdings ('HDFCEQ' etc.); sheet 1: derivatives ('Derivative*').
- Row 1: scheme name banner.
- Row 2: "Portfolio as on DD-MMM-YYYY".
- Row 5: column headers — ISIN | Coupon (%) | Name Of the Instrument |
  Industry+ /Rating | Quantity | Market/ Fair Value (Rs. in Lacs.) |
  % to NAV | Yield | ~YTC.
- Row 6+: section banners ("EQUITY & EQUITY RELATED", "(a) Listed / awaiting
  listing on Stock Exchanges", "Equity", etc.) appear with the label in the
  ISIN column and all other columns None.
- Holding rows: ISIN col B, name col D, industry col E, mkt value col G
  (in lakhs), % to NAV col H.

We yield only rows with a real ISIN (skip section markers, debt totals,
cash/CBLO lines that don't have ISINs). Each adapter row carries
instrument_type derived from the most-recent section banner.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import openpyxl

from mfs import paths
from mfs.io.http import download_to, fetch_bytes
from mfs.ingest.holdings._base import HoldingsAdapter
from mfs.ingest.holdings._registry import register_adapter
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_DISCLOSURE_PAGE = (
    "https://www.hdfcfund.com/statutory-disclosure/portfolio/monthly-portfolio"
)
_FILES_HOST = "files.hdfcfund.com"

# Matches the full HDFC monthly-portfolio Excel URL. We keep the URL match
# tight (no `"` or whitespace allowed in the matched URL) and parse the
# filename downstream to extract the scheme name + data date. Trying to
# capture the scheme name with `.+?` inside the regex was unstable: the
# disclosures HTML embeds the URL inside a JSON-ish blob, and any non-greedy
# capture leaked across the closing quote into adjacent JSON keys.
_URL_RE = re.compile(
    r'(https://files\.hdfcfund\.com/s3fs-public/(\d{4}-\d{2})/'
    r'Monthly%20[^"\s]+?\.xlsx)'
)

_FILENAME_RE = re.compile(
    r"^Monthly\s+(?P<scheme>.+?)\s+-\s+"
    r"(?P<day>\d{1,2})\s+"
    r"(?P<month>January|February|March|April|May|June|"
    r"July|August|September|October|November|December)\s+"
    r"(?P<year>\d{4})\.xlsx$",
    re.IGNORECASE,
)


def _decode_url_path(raw: str) -> str:
    """URL-decode the path component for filename parsing."""
    from urllib.parse import unquote
    return unquote(raw)


def _publish_ym(data_ym: str) -> str:
    """Data month → publish month (publish = data + 1)."""
    y, m = map(int, data_ym.split("-"))
    if m == 12:
        return f"{y + 1:04d}-01"
    return f"{y:04d}-{m + 1:02d}"


def _classify_section(label: str) -> str | None:
    """Map a section banner to an instrument_type, or None if not a section."""
    if not label:
        return None
    s = label.strip().lower()
    if "equity" in s and "derivative" not in s:
        return "Equity"
    if any(k in s for k in (
        "debt", "money market", "tri-party repo", "treasury",
        "government securities", "bonds", "debentures",
    )):
        return "Debt"
    if "reit" in s or "invit" in s:
        return "REIT/InvIT"
    if any(k in s for k in ("net current asset", "cash", "tbill", "t-bill")):
        return "Cash"
    return None


@register_adapter
class HdfcHoldingsAdapter(HoldingsAdapter):
    amc_slug = "hdfc"
    source_label = "HDFC Mutual Fund"

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Scrape the disclosures HTML, return {printed_scheme_name: url}
        for the requested data month (data-month-end date in filename)."""
        html = fetch_bytes(_DISCLOSURE_PAGE).decode("utf-8", errors="replace")
        publish_ym = _publish_ym(ym)
        out: dict[str, str] = {}
        for m in _URL_RE.finditer(html):
            if m.group(2) != publish_ym:
                continue
            url = m.group(1)
            filename = _decode_url_path(url.rsplit("/", 1)[-1])
            fm = _FILENAME_RE.match(filename)
            if not fm:
                log.warning(
                    "holdings.hdfc.filename_unparsed",
                    filename=filename, url=url,
                )
                continue
            scheme = re.sub(r"\s+", " ", fm.group("scheme")).strip()
            # Multiple entries for the same scheme can appear (rare); keep first.
            out.setdefault(scheme, url)
        log.info(
            "holdings.hdfc.discover",
            ym=ym, publish_ym=publish_ym,
            n_schemes=len(out),
        )
        return out

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Download (cached) a single scheme's Excel."""
        out = paths.holdings_excel_raw(self.amc_slug, ym, scheme_filename)
        if out.exists():
            return out
        return download_to(url, out)

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        # Process the equity sheet (sheet 0). Derivatives sheet is intentionally
        # skipped — Active Share is computed on the cash-equity book, and
        # derivatives don't carry ISINs in the same column convention.
        ws = wb[wb.sheetnames[0]]
        current_section: str = "Equity"  # default if banner missing
        for row in ws.iter_rows(values_only=True):
            # We expect at least 8 columns. Defend against jagged rows.
            cells = list(row) + [None] * max(0, 8 - len(row))
            isin = (cells[1] or "")
            name = cells[3]
            weight = cells[7]
            if not isinstance(isin, str):
                isin = str(isin) if isin is not None else ""
            isin = isin.strip()
            # Section banner row: ISIN col is text (not actual ISIN) and
            # weight col is empty.
            if isin and (weight is None or weight == "") and not _is_isin(isin):
                section = _classify_section(isin)
                if section:
                    current_section = section
                continue
            if not _is_isin(isin):
                continue
            if not isinstance(name, str) or not name.strip():
                continue
            try:
                w = float(weight)
            except (TypeError, ValueError):
                continue
            yield ParsedHoldingRecord(
                scheme_name_printed=scheme_name_printed,
                security_name=name.strip(),
                weight_pct=w,
                isin=isin,
                instrument_type=current_section,
                source_amc=self.amc_slug,
            )


_ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}\d$")


def _is_isin(s: str) -> bool:
    """Strict ISIN check: 2 letters + 9 alphanumeric + 1 check digit."""
    return bool(_ISIN_RE.match(s))
