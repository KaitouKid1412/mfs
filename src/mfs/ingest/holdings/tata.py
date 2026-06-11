"""Tata Mutual Fund monthly portfolio holdings adapter (Phase 5).

Like Nippon (and unlike HDFC/SBI's one-Excel-per-scheme), Tata publishes ONE
consolidated workbook per month covering every Tata MF scheme. The portfolio
disclosures page is a Drupal SPA, but its rendered HTML carries every
monthly-portfolio .xls/.xlsx link as a plain absolute URL, so a single page
scrape gives us the whole catalogue:

    https://www.tatamutualfund.com/schemes-related/portfolio

Each link points at a Drupal-managed file under ``/system/files/{YYYY-MM}/``:

    .../system/files/2026-05/Monthly Portfolio as on 30th April 2026 (1).xlsx
    .../system/files/2026-04/Monthly Portfolio as on 31st March 2026.xlsx
    .../system/files/2025-11/Monthly Portfolio as 31st October 2025.xlsx

CRITICAL: the ``{YYYY-MM}`` folder is the PUBLISH month (data month + 1), but
the human date inside the filename ("30th April 2026") is the DATA month-end.
We therefore key discovery off the date PARSED FROM THE FILENAME, not the
folder, so e.g. the April-2026 DATA file (published in the 2026-05 folder) is
correctly selected for ym='2026-04'. The filename date format drifts over time
("as on 30th April", "as 31st October", "as on - 31st January", with optional
"(1)" / "_0" / "_REVISED" / random-hash suffixes), so we extract the
day/month/year tokens with a tolerant regex after normalising '-'/'_' to
spaces. We require "monthly" in the name and exclude the half-yearly-format
files (a different cadence/layout that re-uses March/September month-ends).

Links are served from ``betacms.tatamutualfund.com``, which 302-redirects to
the canonical ``www.tatamutualfund.com``; we rewrite to the canonical host up
front (our HTTP client follows redirects regardless).

Workbook structure (validated against April-2026 file, ~1.6 MB, 70 sheets):
- Sheet 0 ``Index``: row 0 = "Index" title, row 1 = header
  [A=CLASSIFICATION, B=SCHEME CODE, C=SCHEME NAME, D=HYBRID_SCHEME], rows 2+ =
  one row per scheme (code in col B == that scheme's sheet name). Six trailing
  footnote rows put a bare counter (1..6) in col B and are skipped because the
  "code" isn't a real sheet name.
- Non-scheme sheets (``Index``, ``Dividend History``,
  ``Tata Scheme Risk-o-meter``, ``Dummy``) aren't referenced by the Index and
  are never parsed.
- Per-scheme sheet: scheme banner + risk-o-meter + "as on 30-04-2026" preamble,
  then a standard SEBI header (col B="NAME OF THE INSTRUMENT",
  D="INDUSTRY", E="ISIN CODE", G="MKT VAL(Rs. Lacs)", H="% to NAV" in PERCENT
  units), then section banners ("EQUITY & EQUITY RELATED", "A) LISTED/AWAITING
  LISTING…", "MONEY MARKET INSTRUMENTS", …) interleaved with ISIN rows.

Because the per-scheme sheet IS the canonical SEBI layout, we reuse the
validated generic single-sheet parser (``_parse_one_sheet``) — it auto-detects
the header row, ISIN/name/weight columns, the weight unit, dedupes ISINs, and
stops at GRAND TOTAL. The only bespoke parts are (a) discovery of the monthly
URL and (b) Index reverse-lookup to find the right sheet per scheme. Validated
on Tata Flexi Cap (TMCAPF): 56 equity ISIN rows summing to 93.3%.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path
from urllib.parse import unquote

import openpyxl

from mfs import paths
from mfs.ingest.holdings._generic import GenericHoldingsAdapter, _parse_one_sheet
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to, fetch_bytes
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_DISCLOSURE_PAGE = "https://www.tatamutualfund.com/schemes-related/portfolio"
_CANONICAL_HOST = "https://www.tatamutualfund.com"
# The CMS serves links from this beta host, which 302s to the canonical host.
_BETA_HOST_RE = re.compile(
    r"https?://betacms(?:admin)?\.tatamutualfund\.com", re.IGNORECASE
)

# Any .xls/.xlsx URL under /system/files/. We don't anchor the filename here —
# the human-readable date inside it varies too much — and instead parse the
# data month from the filename downstream.
_URL_RE = re.compile(
    r'https?://[^"\'\s<>]+?/system/files/[^"\'\s<>]+?\.(?:xlsx|xls)',
    re.IGNORECASE,
)

_MONTHS = {
    name.lower(): i
    for i, name in enumerate(
        [
            "January", "February", "March", "April", "May", "June",
            "July", "August", "September", "October", "November", "December",
        ],
        start=1,
    )
}

# Day / Month name / 4-digit year appearing anywhere in a (space-normalised)
# filename. Ordinal suffix (st/nd/rd/th) and stray separators (',', '-') are
# optional so "30th April 2026", "31st October 2025", and "- 31st January
# 2025" all parse.
_DATE_RE = re.compile(
    r"(?P<day>\d{1,2})\s*(?:st|nd|rd|th)?[\s\-,]*"
    r"(?P<month>january|february|march|april|may|june|"
    r"july|august|september|october|november|december)"
    r"[\s\-,]*(?P<year>\d{4})",
    re.IGNORECASE,
)


def _canonicalize(url: str) -> str:
    """Rewrite the beta CMS host to the canonical host (skips the 302 hop)."""
    return _BETA_HOST_RE.sub(_CANONICAL_HOST, url)


def _filename_data_ym(filename: str) -> str | None:
    """Reverse-derive the DATA month (YYYY-MM) from a monthly-portfolio
    filename, or None if it isn't a monthly-portfolio file or has no parseable
    date. The date in the name is the data month-end (April data → '2026-04'),
    independent of the publish-month folder it lives in.
    """
    low = filename.lower()
    if "monthly" not in low:
        return None
    # Exclude the half-yearly-format files — different cadence/layout that
    # re-uses 31-March / 30-September month-ends.
    if "half" in low or "yearly" in low:
        return None
    norm = re.sub(r"[-_]+", " ", filename)
    m = _DATE_RE.search(norm)
    if not m:
        return None
    month = _MONTHS[m.group("month").lower()]
    return f"{int(m.group('year')):04d}-{month:02d}"


def _master_excel_filename(ym: str) -> str:
    """Canonical local-cache filename for the consolidated monthly workbook."""
    return f"tata-monthly-portfolio-{ym}.xlsx"


@register_adapter
class TataHoldingsAdapter(GenericHoldingsAdapter):
    amc_slug = "tata"
    source_label = "Tata Mutual Fund"

    # ------------------------------------------------------------------
    # Discovery
    # ------------------------------------------------------------------

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Find the consolidated monthly workbook for data month ``ym``,
        download it once, and enumerate the per-scheme entries from its Index
        sheet. Every entry maps to the SAME master URL (one workbook for all
        schemes); ``fetch_excel`` short-circuits to the cached file.
        """
        master_url = self._select_master_url(ym)
        if master_url is None:
            log.warning("holdings.tata.no_url_for_ym", ym=ym)
            return {}

        master_path = self._cached_master_path(ym)
        if not master_path.exists():
            download_to(master_url, master_path)

        try:
            wb = openpyxl.load_workbook(
                master_path, read_only=True, data_only=True
            )
        except Exception as e:  # noqa: BLE001 — surface any payload issue
            log.error(
                "holdings.tata.workbook_open_failed",
                ym=ym, path=str(master_path), err=str(e),
            )
            return {}

        out: dict[str, str] = {}
        try:
            if "Index" not in wb.sheetnames:
                log.warning(
                    "holdings.tata.no_index_sheet",
                    ym=ym, sheetnames=wb.sheetnames[:5],
                )
                return {}
            sheet_set = set(wb.sheetnames)
            for _code, name in self._iter_index(wb, sheet_set):
                # First write wins on the rare duplicate name.
                out.setdefault(name, master_url)
        finally:
            wb.close()

        log.info(
            "holdings.tata.discover",
            ym=ym, n_schemes=len(out), master_url=master_url,
        )
        return out

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _select_master_url(self, ym: str) -> str | None:
        """Scrape the disclosures HTML for the monthly workbook whose
        filename-encoded DATA month equals ``ym``. Returns the canonicalized
        absolute URL, or None if not yet published.
        """
        html = fetch_bytes(_DISCLOSURE_PAGE).decode("utf-8", errors="replace")
        seen: set[str] = set()
        for mt in _URL_RE.finditer(html):
            url = mt.group(0)
            if url in seen:
                continue
            seen.add(url)
            filename = unquote(url.rsplit("/", 1)[-1])
            if _filename_data_ym(filename) == ym:
                return _canonicalize(url)
        return None

    def _cached_master_path(self, ym: str) -> Path:
        return paths.holdings_excel_raw(
            self.amc_slug, ym, _master_excel_filename(ym)
        )

    @staticmethod
    def _iter_index(
        wb: openpyxl.Workbook, sheet_set: set[str]
    ) -> Iterable[tuple[str, str]]:
        """Yield (sheet_code, scheme_name) for every Index row whose SCHEME
        CODE (col B) is a real sheet. Skips the title/header rows and the
        trailing footnote rows (whose col-B "code" isn't a sheet name).
        """
        idx_ws = wb["Index"]
        for i, row in enumerate(idx_ws.iter_rows(values_only=True)):
            if i < 2:  # row 0 = "Index" title, row 1 = column header
                continue
            code = row[1] if len(row) > 1 else None
            name = row[2] if len(row) > 2 else None
            if not code or not isinstance(name, str):
                continue
            code_s = str(code).strip()
            if code_s not in sheet_set:
                continue
            name_s = re.sub(r"\s+", " ", name).strip()
            if not name_s:
                continue
            yield code_s, name_s

    # ------------------------------------------------------------------
    # Download
    # ------------------------------------------------------------------

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Return the cached master workbook for this month. The URL is shared
        across all schemes, so ``scheme_filename`` is ignored. We still
        download (idempotent) in case fetch is called without discover first.
        """
        master_path = self._cached_master_path(ym)
        if master_path.exists():
            return master_path
        return download_to(_canonicalize(url), master_path)

    # ------------------------------------------------------------------
    # Parse
    # ------------------------------------------------------------------

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        """Locate the scheme's sheet via the Index, then delegate to the
        validated generic single-sheet SEBI parser.
        """
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            sheet_code = self._find_sheet_for_scheme(wb, scheme_name_printed)
            if sheet_code is None:
                log.warning(
                    "holdings.tata.scheme_not_in_index",
                    scheme=scheme_name_printed, ym=ym,
                )
                return
            recs = _parse_one_sheet(wb[sheet_code])
            if not recs:
                log.warning(
                    "holdings.tata.empty_sheet",
                    scheme=scheme_name_printed, code=sheet_code, ym=ym,
                )
                return
            for r in recs:
                yield ParsedHoldingRecord(
                    scheme_name_printed=scheme_name_printed,
                    security_name=r.security_name,
                    weight_pct=r.weight_pct,
                    isin=r.isin,
                    instrument_type=r.instrument_type,
                    source_amc=self.amc_slug,
                )
        finally:
            wb.close()

    def _find_sheet_for_scheme(
        self, wb: openpyxl.Workbook, scheme_name_printed: str
    ) -> str | None:
        """Reverse-lookup the scheme's sheet code from the Index sheet,
        case-insensitively on the printed scheme name."""
        if "Index" not in wb.sheetnames:
            return None
        target = re.sub(r"\s+", " ", scheme_name_printed.strip()).lower()
        sheet_set = set(wb.sheetnames)
        for code, name in self._iter_index(wb, sheet_set):
            if name.lower() == target:
                return code
        return None
