"""Sundaram Mutual Fund monthly portfolio holdings adapter (Phase 5).

Sundaram publishes TWO consolidated workbooks per month — one "Fixed
Income" and one "Equity & Fund of Funds" — each covering every scheme of
that family in a single .xlsx with one sheet per scheme (like Nippon, not
HDFC's one-file-per-scheme). We only need the Equity & Fund of Funds
workbook for the ranked equity universe.

Discovery (SPA + AjaxPro backing call)
--------------------------------------
The portfolios page ``/monthly-fortnightly-adhoc-portfolios`` is an
ASP.NET/AjaxPro SPA: the file list is injected client-side by calling a
server method. We reproduce that XHR directly:

    POST /ajax/Modules_Disclosure_Monthly_Fortnightly_Adhoc_Portfolios,
         App_Web_prf1m21b.ashx?_method=GetCategory&_session=no
    body: Catid=Monthly

The ``App_Web_*`` token in that path is a server-build hash that can change
when Sundaram redeploys, so we don't hardcode it: we GET the page first and
scrape the actual ``<script src=".../...GetCategory...ashx">`` URL out of
its HTML. The response is a single-quoted JS string literal whose inner
quotes are backslash-escaped; we unescape it, then it's plain HTML carrying
one anchor per published workbook:

    <a href='/uploaddir/MonthlyPortfolio/monthlyportfolio_090526110429.xlsx'
       ...>Monthly Portfolio Disclosure Equity & Fund of Funds - Apr 2026</a>

The href filename is an opaque upload-timestamp, so we key off the anchor
TITLE instead: ``Monthly Portfolio Disclosure Equity & Fund of Funds -
<Mon> <YYYY>`` where ``<Mon>`` is a 3-letter English month abbreviation.
For data month ``2026-04`` we select the title containing ``Apr 2026``.

Workbook structure (validated against the April 2026 Equity & FoF file)
----------------------------------------------------------------------
- Sheet 0 == ``Index``: col A = S.NO., col B = ACRONYM (== the per-scheme
  sheet name), col C = full SCHEME NAME.
- Sheets 1..N: one per scheme, sheet name == the Index acronym.
- Per-scheme layout:
    row 0: 'SUNDARAM MUTUAL FUND' banner
    row 1: scheme name
    row 2: 'Monthly Portfolio Statement for ...'
    row 3: header — col B='ISIN Code', col C='Name of the instrument',
           col D='Rating / Industry', col E='Quantity',
           col F='Mkt Value Rs. in Lacs', col G='% of Net Asset', col H='YTM (%)'
    row 4+: section banners ('A) Equity & Equity Related',
            '(a) Listed / awaiting listing on ...') + holding rows.
- CRITICAL UNIT NOTE: like Nippon, '% of Net Asset' is stored as a
  **fraction** (0.03346 == 3.35%), not a percentage.

Parsing
-------
The shared ``parse_sebi_excel`` already detects this exact layout — header
row, ISIN/name/weight columns, and the fraction-vs-percent unit (it scales
a sheet whose weights sum to <= 1.5 by 100). So ``parse_excel`` just
resolves the scheme's sheet index from the Index sheet and delegates one
sheet to ``parse_sebi_excel`` (modeled on the per-sheet approach in
``nippon.py``).
"""

from __future__ import annotations

import html as _html
import re
from collections.abc import Iterable
from pathlib import Path
from urllib.parse import urljoin

import httpx
import openpyxl

from mfs import paths
from mfs.ingest.holdings._generic import GenericHoldingsAdapter, parse_sebi_excel
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to, fetch_bytes
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_PORTFOLIOS_PAGE = (
    "https://www.sundarammutual.com/monthly-fortnightly-adhoc-portfolios"
)
_BASE = "https://www.sundarammutual.com"

_UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"
)

# The AjaxPro proxy <script> for the GetCategory server method. The
# App_Web_* build hash is volatile, so we discover the live path from the
# page HTML rather than hardcoding it.
_GETCATEGORY_SCRIPT_RE = re.compile(
    r'src="(?P<url>/ajax/Modules_Disclosure_Monthly_Fortnightly_Adhoc_'
    r'Portfolios[^"]*?\.ashx)"',
    re.IGNORECASE,
)

# One workbook per family per month. We only want "Equity & Fund of Funds".
# Title shape: "Monthly Portfolio Disclosure Equity & Fund of Funds - <Mon> <YYYY>".
# Month is a 3-letter English abbreviation; spacing around the dash varies.
_MONTH_ABBR = (
    "Jan", "Feb", "Mar", "Apr", "May", "Jun",
    "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
)

# The anchor wraps a leading <i> icon tag before the visible title text:
#   <a href='...xlsx' ...><i class='...'></i>Monthly Portfolio Disclosure
#   Equity & Fund of Funds - Apr 2026</a>
# So after the opening tag we skip any number of nested tags, then capture
# the trailing text node as the title.
_ANCHOR_RE = re.compile(
    r"href='(?P<url>/uploaddir/MonthlyPortfolio/[^']+?\.xlsx)'"
    r"[^>]*>(?:<[^>]+>)*"
    r"(?P<title>[^<]*?Equity\s*&\s*Fund of Funds[^<]*?)</a>",
    re.IGNORECASE,
)


def _data_month_label(ym: str) -> str:
    """data_ym='2026-04' -> 'Apr 2026' (the token in the anchor title)."""
    y, m = map(int, ym.split("-"))
    return f"{_MONTH_ABBR[m - 1]} {y:04d}"


# Sundaram's Index sheet spells two flagship funds differently from AMFI's
# scheme_master recorded name. The shared fuzzy matcher (token_set_ratio +
# Levenshtein tie-break) mis-resolves the workbook spellings:
#   - "Sundaram Flexi Cap Fund"        -> mis-matched to "Sundaram Mid Cap
#     Fund" (the split "Flexi Cap" loses the single-token "FLEXICAP" signal).
#   - "Sundaram Large And Mid Cap Fund" -> mis-matched to "Sundaram Large Cap
#     Fund" at a perfect token_set_ratio of 100 (the shorter candidate is a
#     pure token-subset, and "MID CAP" != master's single-token "MIDCAP").
# We can't touch the shared matcher; instead we normalize the printed name we
# emit so it lands on the master's spelling. This is a pure rename (no value
# fabrication) and is applied symmetrically in discovery and parse so the
# Index sheet lookup still resolves.
_NAME_NORMALIZE = {
    "sundaram flexi cap fund": "Sundaram Flexicap Fund",
    "sundaram large and mid cap fund": "Sundaram Large and Midcap Fund",
}


def _normalize_index_name(name: str) -> str:
    """Map a workbook Index scheme name to the AMFI-spelled name when the two
    diverge in a way that breaks the shared fuzzy matcher; otherwise return
    the name unchanged."""
    key = re.sub(r"\s+", " ", name.replace("\xa0", " ")).strip().lower()
    return _NAME_NORMALIZE.get(key, name)


def _unescape_ajaxpro(body: str) -> str:
    """The AjaxPro response is a JS string literal: a leading/trailing
    single quote with inner quotes backslash-escaped (\\' and \\"). Strip
    the wrapping quotes and unescape so it's plain HTML we can regex."""
    s = body.strip()
    if len(s) >= 2 and s[0] == "'" and s[-1] == "'":
        s = s[1:-1]
    return s.replace("\\'", "'").replace('\\"', '"').replace("\\/", "/")


@register_adapter
class SundaramHoldingsAdapter(GenericHoldingsAdapter):
    amc_slug = "sundaram"
    source_label = "Sundaram Mutual Fund"

    # ------------------------------------------------------------------
    # Discovery
    # ------------------------------------------------------------------

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Find the Equity & FoF consolidated workbook for `ym`, download it
        once, and map every scheme name (from its Index sheet) to that URL.

        Like Nippon, Sundaram ships ONE workbook for all (equity) schemes,
        so every entry in the returned dict points at the same URL.
        """
        master_url = self._discover_master_url(ym)
        if master_url is None:
            log.warning("holdings.sundaram.no_url_for_ym", ym=ym)
            return {}

        master_path = self._cached_master_path(ym)
        if not master_path.exists():
            download_to(master_url, master_path)

        out: dict[str, str] = {}
        try:
            wb = openpyxl.load_workbook(
                master_path, read_only=True, data_only=True
            )
        except Exception as e:  # noqa: BLE001 — surface any payload issue
            log.error(
                "holdings.sundaram.workbook_open_failed",
                ym=ym, path=str(master_path), err=str(e),
            )
            return {}
        try:
            for _code, name in self._index_entries(wb):
                out.setdefault(_normalize_index_name(name), master_url)
        finally:
            wb.close()

        log.info(
            "holdings.sundaram.discover",
            ym=ym, n_schemes=len(out), master_url=master_url,
        )
        return out

    def _discover_master_url(self, ym: str) -> str | None:
        """Reproduce the page's AjaxPro XHR and return the absolute URL of
        the Equity & FoF workbook whose title matches data month `ym`."""
        page = fetch_bytes(_PORTFOLIOS_PAGE).decode("utf-8", errors="replace")
        m = _GETCATEGORY_SCRIPT_RE.search(page)
        if not m:
            log.error("holdings.sundaram.no_getcategory_script", ym=ym)
            return None
        ashx = urljoin(_BASE + "/", m.group("url").lstrip("/"))
        endpoint = f"{ashx}?_method=GetCategory&_session=no"

        with httpx.Client(
            timeout=60.0,
            headers={
                "User-Agent": _UA,
                "X-AjaxPro-Method": "GetCategory",
                "X-Requested-With": "XMLHttpRequest",
                "Content-Type": "text/plain; charset=UTF-8",
                "Referer": _PORTFOLIOS_PAGE,
                "Accept": "*/*",
            },
            follow_redirects=True,
        ) as c:
            r = c.post(endpoint, content=b"Catid=Monthly")
            r.raise_for_status()
            body = r.text

        html = _unescape_ajaxpro(body)
        want = _data_month_label(ym).lower()
        for am in _ANCHOR_RE.finditer(html):
            title = re.sub(r"\s+", " ", _html.unescape(am.group("title"))).strip()
            # Title ends with "- <Mon> <YYYY>"; require the exact month label.
            if title.lower().endswith(want):
                url = am.group("url")
                return urljoin(_BASE + "/", url.lstrip("/"))
        log.warning("holdings.sundaram.no_anchor_for_ym", ym=ym, want=want)
        return None

    @staticmethod
    def _index_entries(wb: openpyxl.Workbook) -> list[tuple[str, str]]:
        """Yield (sheet_code, scheme_name) for each Index row whose code
        names an actual sheet in the workbook."""
        if "Index" not in wb.sheetnames:
            return []
        sheet_set = set(wb.sheetnames)
        out: list[tuple[str, str]] = []
        for i, row in enumerate(wb["Index"].iter_rows(values_only=True)):
            if i == 0:
                continue  # 'S.NO. | ACRONYM | SCHEME NAME' header
            code = row[1] if len(row) > 1 else None
            name = row[2] if len(row) > 2 else None
            if not code or not isinstance(name, str):
                continue
            code_s = str(code).strip()
            name_s = re.sub(r"\s+", " ", name.replace("\xa0", " ")).strip()
            if not code_s or not name_s or code_s not in sheet_set:
                continue
            out.append((code_s, name_s))
        return out

    def _cached_master_path(self, ym: str) -> Path:
        return paths.holdings_excel_raw(
            self.amc_slug, ym, f"sundaram-equity-fof-{ym}.xlsx"
        )

    # ------------------------------------------------------------------
    # Download (shared workbook → ignore the per-scheme filename)
    # ------------------------------------------------------------------

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        master_path = self._cached_master_path(ym)
        if master_path.exists():
            return master_path
        return download_to(url, master_path)

    # ------------------------------------------------------------------
    # Parse (resolve the scheme's sheet, delegate to the generic parser)
    # ------------------------------------------------------------------

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            target = re.sub(
                r"\s+", " ", scheme_name_printed.replace("\xa0", " ")
            ).strip().lower()
            sheet_idx: int | None = None
            for code, name in self._index_entries(wb):
                # Compare on the normalized (AMFI-spelled) form so the
                # rewritten flagship names from discovery still resolve to
                # their original Index sheet.
                if _normalize_index_name(name).lower() == target:
                    sheet_idx = wb.sheetnames.index(code)
                    break
        finally:
            wb.close()

        if sheet_idx is None:
            log.warning(
                "holdings.sundaram.scheme_not_in_index",
                scheme=scheme_name_printed, ym=ym,
            )
            return

        # The generic parser detects the header row, columns, and the
        # fraction weight unit on this single sheet.
        yield from parse_sebi_excel(
            excel_path, scheme_name_printed, self.amc_slug,
            sheet_index=sheet_idx,
        )
