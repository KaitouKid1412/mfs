"""UTI Mutual Fund monthly portfolio holdings adapter (Phase 5).

Like Nippon (and unlike HDFC/SBI's per-scheme files), UTI publishes ONE
consolidated SEBI-format workbook covering every UTI scheme each month. It is
shipped inside a ZIP on UTI's CloudFront CDN, surfaced from the official
"Consolidate All Portfolio Disclosure" page
(https://www.utimf.com/downloads/consolidate-all-portfolio-disclosure):

    https://d3ce1o48hc5oli.cloudfront.net/s3fs-public/{PUBLISH_YM}/
        uti_mf_scheme_portfolios_{DD.MM.YYYY}.zip

where:
- PUBLISH_YM  = data month + 1 (April-2026 data lands in the 2026-05 folder,
  same publish-lag convention as HDFC/SBI/Nippon).
- DD.MM.YYYY  = the data-month-END date (30.04.2026 for April-2026 data).

The ZIP carries four files; the one we want is the consolidated portfolio
workbook named ``Sebi Exposure <DD> <Mon> <YYYY>_final.xlsx`` (a few legacy
months ship it as ``.xls`` — we match on the "sebi exposure" stem and the
``.xls``/``.xlsx`` suffix, ignoring the Divmast / Futures-exposure /
Risk-o-meter siblings). Despite the "Sebi Exposure" name it is the full
ISIN-bearing monthly portfolio statement, not a derivatives-only file.

Consolidated-workbook layout (validated against April-2026, 8.3k rows):
- A SINGLE sheet named ``exposure`` holds every scheme back-to-back.
- Each scheme block is delimited by marker rows in column A:
    ``SCHEME CODE<nnn>STARTS``  …block…  ``SCHEME CODE<nnn>ENDS``
- Inside a block:
    row +0  : ``SCHEME CODE<nnn>STARTS``
    row +1  : ``UTI MUTUAL FUND``
    row +2  : ``SCHEME: <printed scheme name>``  ← the name we expose
    row +3  : ``PROVISIONAL AND UNAUDITED PORTFOLIO DISCLOSURE AS OF …``
    row +5  : header  (col A=NAME, col B=RATING/INDUSTRY, col C=QUANTITY,
              col D=MARKET-VALUE, col E=% TO NAV, col H=ISIN)
    row +6+ : section banners ("EQUITY AND EQUITY RELATED",
              "MONEY MARKET INSTRUMENTS", "NET CURRENT ASSETS", …) and
              holding rows, then "TOTAL …" footers and a per-scheme tail of
              notes, then ``SCHEME CODE<nnn>ENDS``.
- Holding rows: name col A (prefixed "EQ - ", "DB - ", etc.), industry col B,
  ISIN col H, ``% TO NAV`` col E. ``% TO NAV`` is stored as a PERCENT (2.66 =
  2.66%), not a fraction, so no scaling is needed.

Discovery + fetch follow the Nippon model: every scheme maps to the same ZIP
URL, ``fetch_excel`` short-circuits to the already-extracted cached workbook,
and ``parse_excel`` walks just the requested scheme's block.
"""

from __future__ import annotations

import calendar
import re
import zipfile
from collections.abc import Iterable
from pathlib import Path

import openpyxl

from mfs import paths
from mfs.ingest.holdings._generic import (
    GenericHoldingsAdapter,
    classify_section,
    is_isin,
)
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_CDN_HOST = "https://d3ce1o48hc5oli.cloudfront.net"

# Column offsets inside the consolidated "exposure" sheet (0-indexed).
_COL_NAME = 0
_COL_ISIN = 7
_COL_WEIGHT = 4

# A scheme block opens with "SCHEME CODE<nnn>STARTS" in column A and the
# printed name appears on a later "SCHEME: <name>" row within the block.
_STARTS_RE = re.compile(r"^SCHEME\s+CODE\w+STARTS$", re.IGNORECASE)
_ENDS_RE = re.compile(r"^SCHEME\s+CODE\w+ENDS$", re.IGNORECASE)
_SCHEME_NAME_RE = re.compile(r"^SCHEME:\s*(?P<name>.+?)\s*$", re.IGNORECASE)

# Inner-zip member we want: the consolidated portfolio workbook. A handful of
# legacy months ship it as .xls; the file is OOXML either way but openpyxl
# refuses the .xls suffix, so we always cache it as .xlsx.
_PORTFOLIO_MEMBER_RE = re.compile(r"sebi\s+exposure\b.*\.(xlsx|xls)$", re.IGNORECASE)


def _publish_ym(data_ym: str) -> str:
    """Data month → publish month (UTI publishes data+1, like HDFC/SBI)."""
    y, m = map(int, data_ym.split("-"))
    if m == 12:
        return f"{y + 1:04d}-01"
    return f"{y:04d}-{m + 1:02d}"


def _data_month_end(data_ym: str) -> str:
    """data_ym='2026-04' → '30.04.2026' (DD.MM.YYYY of the month-end)."""
    y, m = map(int, data_ym.split("-"))
    dd = calendar.monthrange(y, m)[1]
    return f"{dd:02d}.{m:02d}.{y:04d}"


def _zip_url(data_ym: str) -> str:
    """Build the canonical CloudFront ZIP URL for a data month."""
    return (
        f"{_CDN_HOST}/s3fs-public/{_publish_ym(data_ym)}/"
        f"uti_mf_scheme_portfolios_{_data_month_end(data_ym)}.zip"
    )


def _norm_name(name: str) -> str:
    """Normalize a scheme name for matching: lowercase, collapse whitespace,
    drop a trailing period. UTI prints e.g. 'UTI - Flexi Cap Fund.' with a
    stray trailing dot that the scheme_master copy lacks."""
    s = re.sub(r"\s+", " ", name).strip().lower()
    return s.rstrip(".").strip()


@register_adapter
class UtiHoldingsAdapter(GenericHoldingsAdapter):
    amc_slug = "uti"
    source_label = "UTI Mutual Fund"

    # ------------------------------------------------------------------
    # Cache helpers
    # ------------------------------------------------------------------

    def _cached_master_path(self, ym: str) -> Path:
        """Path for the extracted consolidated workbook for this month."""
        return paths.holdings_excel_raw(
            self.amc_slug, ym, f"uti_consolidated_{ym}.xlsx"
        )

    def _ensure_master(self, ym: str) -> Path:
        """Download the month's ZIP (if needed), extract the consolidated
        portfolio workbook, and return the cached local path."""
        master_path = self._cached_master_path(ym)
        if master_path.exists():
            return master_path

        # Download the ZIP to a sibling temp path under the same dir.
        zip_path = paths.holdings_excel_raw(
            self.amc_slug, ym, f"uti_scheme_portfolios_{ym}.zip"
        )
        if not zip_path.exists():
            download_to(_zip_url(ym), zip_path)

        with zipfile.ZipFile(zip_path) as zf:
            member = next(
                (n for n in zf.namelist() if _PORTFOLIO_MEMBER_RE.search(n)),
                None,
            )
            if member is None:
                raise RuntimeError(
                    f"UTI ZIP for {ym} has no 'Sebi Exposure' portfolio "
                    f"workbook; members={zf.namelist()}"
                )
            data = zf.read(member)

        master_path.parent.mkdir(parents=True, exist_ok=True)
        master_path.write_bytes(data)
        return master_path

    # ------------------------------------------------------------------
    # Discovery
    # ------------------------------------------------------------------

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Return {printed_scheme_name: zip_url} for the data month `ym`.

        UTI ships ONE consolidated workbook, so every scheme maps to the same
        ZIP URL (Nippon-style). We open the workbook and enumerate the
        'SCHEME: <name>' rows that follow each 'SCHEME CODE…STARTS' marker.
        """
        master_path = self._ensure_master(ym)
        url = _zip_url(ym)

        wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
        out: dict[str, str] = {}
        try:
            ws = self._portfolio_sheet(wb)
            in_block = False
            for row in ws.iter_rows(values_only=True):
                c0 = row[0] if row else None
                if not isinstance(c0, str):
                    continue
                s = c0.strip()
                if _STARTS_RE.match(s):
                    in_block = True
                    continue
                if _ENDS_RE.match(s):
                    in_block = False
                    continue
                if not in_block:
                    continue
                m = _SCHEME_NAME_RE.match(s)
                if m:
                    name = re.sub(r"\s+", " ", m.group("name")).strip()
                    if name:
                        out.setdefault(name, url)
                    in_block = False  # name found; ignore rest of block here
        finally:
            wb.close()

        log.info("holdings.uti.discover", ym=ym, n_schemes=len(out), zip_url=url)
        return out

    # ------------------------------------------------------------------
    # Fetch
    # ------------------------------------------------------------------

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Return the cached consolidated workbook for this month.

        UTI's workbook is shared across all schemes, so `url` /
        `scheme_filename` are intentionally ignored — we ensure the ZIP is
        downloaded + extracted and hand back the single local workbook.
        """
        return self._ensure_master(ym)

    # ------------------------------------------------------------------
    # Parse
    # ------------------------------------------------------------------

    @staticmethod
    def _portfolio_sheet(wb: openpyxl.Workbook):
        """The consolidated portfolio lives on the 'exposure' sheet; fall back
        to sheet 0 if UTI ever renames it."""
        for name in wb.sheetnames:
            if name.strip().lower() == "exposure":
                return wb[name]
        return wb.worksheets[0]

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        """Walk just the requested scheme's block and yield ISIN-bearing rows.

        We locate the block whose 'SCHEME: <name>' matches (normalized),
        track the most-recent section banner for instrument_type, skip
        TOTAL/footer rows, and stop at the block's ENDS marker. '% TO NAV' is
        already a percentage, so weights pass through unscaled.
        """
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            ws = self._portfolio_sheet(wb)
            target = _norm_name(scheme_name_printed)

            in_target = False  # inside the matched scheme's block
            scanning_block = False  # inside SOME block, name not yet checked
            current_section = "Equity"
            seen: set[str] = set()

            for row in ws.iter_rows(values_only=True):
                cells = list(row) + [None] * max(0, _COL_ISIN + 1 - len(row))
                c0 = cells[_COL_NAME]
                c0s = c0.strip() if isinstance(c0, str) else ""

                if _STARTS_RE.match(c0s):
                    scanning_block = True
                    in_target = False
                    current_section = "Equity"
                    continue
                if _ENDS_RE.match(c0s):
                    if in_target:
                        break  # finished the scheme we wanted
                    scanning_block = False
                    continue

                if scanning_block and not in_target:
                    m = _SCHEME_NAME_RE.match(c0s)
                    if m:
                        in_target = _norm_name(m.group("name")) == target
                    continue

                if not in_target:
                    continue

                isin = cells[_COL_ISIN]
                weight = cells[_COL_WEIGHT]
                isin_s = str(isin).strip() if isin is not None else ""

                # Section banner / TOTAL footer: a non-ISIN label row.
                if c0s and not is_isin(isin_s):
                    section = classify_section(c0s)
                    if section:
                        current_section = section
                    continue

                if not is_isin(isin_s) or not c0s:
                    continue
                try:
                    w = float(weight)
                except (TypeError, ValueError):
                    continue
                if isin_s in seen:
                    continue
                seen.add(isin_s)

                # UTI prefixes names with a 2-char instrument tag, e.g.
                # "EQ - HDFC BANK LIMITED" / "DB - ...". Strip it for a clean
                # security_name; keep everything after the first " - ".
                name = re.sub(r"^[A-Z]{2,4}\s*-\s*", "", c0s).strip()

                yield ParsedHoldingRecord(
                    scheme_name_printed=scheme_name_printed,
                    security_name=name or c0s,
                    weight_pct=w,
                    isin=isin_s,
                    instrument_type=current_section,
                    source_amc=self.amc_slug,
                )
        finally:
            wb.close()
