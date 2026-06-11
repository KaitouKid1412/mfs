"""Motilal Oswal Mutual Fund monthly portfolio holdings adapter (Phase 5).

Like Nippon (and unlike HDFC/SBI's one-Excel-per-scheme model), Motilal
Oswal publishes ONE consolidated workbook per month covering every MOMF
scheme. Discovery is a JSON-API problem, not an HTML-scrape problem: the
downloads page (``/downloads/scheme-portfolio-details``) is an Adobe Edge
Delivery SPA whose document table is populated client-side from an AEM
backend search endpoint:

    GET https://www.motilaloswalmf.com/content/aem-cloud-dept-backend-
        motilal-oswal/api/search-documents.json
        ?year=<YYYY>&category=month%20end%20portfolio&month=<mon>&type=mf

…returning ``{"results":[{"title","path","publishDate","month","year"}]}``
where ``path`` is a DAM path under
``/content/dam/motilal-mf/downloads/mf/month-end-portfolio/<year>/<mon>/``.

The category string ``month end portfolio`` and ``type=mf`` were recovered
from the block JS (``/blocks/our-funds-block/our-funds-block.js`` →
``dataMapMoObj.keyval['Scheme Portfolio Details']`` and the section-derived
``param.type``). ``month`` is the 3-letter lowercase publish month.

PUBLISH vs DATA month: the monthly portfolio "as on 30-Apr-2026" (data
month 2026-04) is published in MAY 2026, so it lands in the ``may`` folder
under a title ``scheme portfolio details - april 2026``. The DATA month is
carried in the TITLE, not the folder. So discovery:
  1. query the publish-month folder (data month + 1) — the normal home,
  2. fall back to the data-month folder (in case MO files it early),
  and in both cases keep only the result whose TITLE is the canonical
  ``scheme portfolio details - <data-month-name> <year>`` monthly file
  (NOT the ``fortnightly``/``half yearly`` siblings in the same folder).

Workbook structure (validated against the April-2026 file, 85 sheets):
- Sheet ``INDEX``: Fund Name in col D (idx 3), 4-char Fund Code in col E
  (idx 4), one row per scheme; codes look like ``YO08``.
- Sheets ``YOnn``: one per scheme, sheet name == the Fund Code.
- Per-scheme layout is a clean SEBI portfolio table — header row carries
  ``Name of the Instrument`` | ``ISIN Code`` | ``Industry
  Classification*`` | ``Quantity`` | ``Market Value (Rs. in Lakhs)`` |
  ``% to NAV``, with ``(A) Equity & Equity related`` / ``(a) Listed /
  awaiting listing on Stock Exchanges`` section banners and ``% to NAV``
  printed as a true percent (7.36 == 7.36%). Because the layout is
  generic-SEBI, we reuse the shared ``_parse_one_sheet`` (auto-detects the
  header row, ISIN/name/weight columns, and the percent-vs-fraction unit)
  rather than hard-coding column offsets.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path

import openpyxl

from mfs import paths
from mfs.ingest.holdings._generic import GenericHoldingsAdapter, _parse_one_sheet
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to, fetch_bytes
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_SEARCH_API = (
    "https://www.motilaloswalmf.com/content/"
    "aem-cloud-dept-backend-motilal-oswal/api/search-documents.json"
)
_HOST = "https://www.motilaloswalmf.com"

# Category + type identifiers, recovered from the SPA block JS.
_CATEGORY = "month end portfolio"
_TYPE = "mf"

_MONTHS_SHORT = [
    "jan", "feb", "mar", "apr", "may", "jun",
    "jul", "aug", "sep", "oct", "nov", "dec",
]
_MONTHS_FULL = [
    "january", "february", "march", "april", "may", "june",
    "july", "august", "september", "october", "november", "december",
]


def _publish_ym(data_ym: str) -> str:
    """Data month -> publish month (MO publishes the monthly portfolio in
    the following calendar month, so the file lands in the data+1 folder)."""
    y, m = map(int, data_ym.split("-"))
    if m == 12:
        return f"{y + 1:04d}-01"
    return f"{y:04d}-{m + 1:02d}"


def _ym_parts(ym: str) -> tuple[int, str, str]:
    """ym='2026-04' -> (2026, 'apr', 'april'). Folder month is the 3-letter
    lowercase short name; the canonical title uses the full month name."""
    y, m = map(int, ym.split("-"))
    return y, _MONTHS_SHORT[m - 1], _MONTHS_FULL[m - 1]


# The canonical monthly file's title is "scheme portfolio details -
# <full-month> <year>" (case-insensitive, tolerant of stray leading/inner
# whitespace MO occasionally injects). This deliberately EXCLUDES the
# "fortnightly portfolio report …" and "half yearly portfolio …" siblings
# that share the same folder, and the factsheet-derived
# "IN_MF_..._FACTSHEET" entries.
def _is_monthly_title(title: str, data_ym: str) -> bool:
    _, _, full = _ym_parts(data_ym)
    year = data_ym.split("-")[0]
    norm = re.sub(r"\s+", " ", title.strip().lower())
    return norm == f"scheme portfolio details - {full} {year}"


def _consolidated_filename(ym: str) -> str:
    """Local-cache filename for the consolidated monthly workbook."""
    return f"motilal-oswal-monthly-portfolio-{ym}.xlsx"


def _query_folder(year: int, month_short: str) -> list[dict]:
    """Call the AEM search-documents API for one (year, publish-month)
    folder and return its ``results`` list (empty on any failure)."""
    params = {
        "year": str(year),
        "category": _CATEGORY,
        "month": month_short,
        "type": _TYPE,
    }
    try:
        raw = fetch_bytes(_SEARCH_API, params=params)
    except Exception as e:  # noqa: BLE001 — surface as empty + log
        log.warning(
            "holdings.motilal_oswal.search_failed",
            year=year, month=month_short, err=str(e),
        )
        return []
    import json

    try:
        data = json.loads(raw.decode("utf-8", errors="replace"))
    except json.JSONDecodeError as e:
        log.warning(
            "holdings.motilal_oswal.search_bad_json",
            year=year, month=month_short, err=str(e),
        )
        return []
    results = data.get("results")
    return results if isinstance(results, list) else []


def _extract_scheme_name(raw: object) -> str | None:
    """Clean a scheme name from the INDEX sheet's Fund-Name cell.

    Collapses whitespace and drops the trailing parenthesised
    "(Formerly known as …)" / SEBI-category description so the printed name
    matches scheme_master. Returns None for empty / non-string cells.
    """
    if not isinstance(raw, str):
        return None
    # NBSP and other Unicode spaces appear inside MO's Fund-Name cells
    # (e.g. "Motilal Oswal\xa0Services Fund"); \s in re collapses them too.
    name = re.sub(r"\s+", " ", raw).strip()
    if not name:
        return None
    # Strip a trailing parenthetical (formerly-known-as / category blurb).
    name = re.sub(r"\s*\([^()]*\)\s*$", "", name).strip()
    return name or None


def _resolve_sheet_code(raw_code: object, sheet_set: set[str]) -> str | None:
    """Map an INDEX Fund-Code cell to the workbook sheet it names.

    MO's per-scheme sheets are named ``YOnn`` (letter-O), but the INDEX
    sheet occasionally types the code with a digit-zero instead of the
    letter (observed in the April-2026 file: Business Cycle's sheet is
    ``YO54`` yet INDEX prints ``Y054``). When the literal code is not a
    sheet name, retry with the leading ``Y0`` rewritten to ``YO`` — the
    only OCR/typo variant seen — and accept it only if THAT names a real
    sheet. Returns the resolved sheet-code string, or None if neither
    form names a sheet (so non-scheme rows like the ``Fund Code`` header
    are still skipped).
    """
    if not raw_code:
        return None
    code = str(raw_code).strip()
    if code in sheet_set:
        return code
    if code.startswith("Y0"):
        alt = "YO" + code[2:]
        if alt in sheet_set:
            return alt
    return None


# Printed-name disambiguation for funds whose cleaned INDEX name is a
# strict token-SUPERSET of a shorter MO sibling. The shared scheme matcher
# scores both at token_set_ratio 100 and breaks the tie on Levenshtein
# ``ratio``; when the longer fund's scheme_master key carries an un-stripped
# "DIRECT GROWTH" suffix (AMFI names it "Fund- Direct Growth", which the
# matcher's "- Direct PLAN" suffix rule does not strip), the shorter sibling
# wins and the longer fund is silently misrouted to it.
#
# Concretely: INDEX "Motilal Oswal Financial Services Fund" was matched to
# scheme_master "Motilal Oswal Services Fund" (153558), stealing the real
# Financial Services Fund's (154142) holdings. We re-emit the longer fund's
# name with the literal "- Direct Growth" plan/option tokens that scheme_master
# keeps in its canonical key, so the input becomes an exact-length match and
# wins the tie-break. The added tokens are the fund's true plan/option, not a
# fabricated data value, and ``_find_sheet_for_scheme`` strips them back off
# before the INDEX sheet lookup. Keyed on the cleaned (whitespace-collapsed)
# INDEX name; only applied when that name is found verbatim.
_DISAMBIGUATION_SUFFIX = "- Direct Growth"
_NEEDS_DISAMBIGUATION = frozenset({
    "Motilal Oswal Financial Services Fund",
})


def _emit_name(cleaned: str) -> str:
    """Apply printed-name disambiguation for collision-prone MO funds."""
    if cleaned in _NEEDS_DISAMBIGUATION:
        return f"{cleaned} {_DISAMBIGUATION_SUFFIX}"
    return cleaned


def _strip_disambiguation(printed: str) -> str:
    """Reverse :func:`_emit_name` for INDEX sheet re-lookup."""
    suffix = f" {_DISAMBIGUATION_SUFFIX}"
    if printed.endswith(suffix):
        return printed[: -len(suffix)]
    return printed


@register_adapter
class MotilalOswalHoldingsAdapter(GenericHoldingsAdapter):
    amc_slug = "motilal_oswal"
    source_label = "Motilal Oswal Mutual Fund"

    # ------------------------------------------------------------------
    # Discovery
    # ------------------------------------------------------------------

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Locate the consolidated monthly workbook for data month ``ym``,
        download it once, and enumerate per-scheme entries from its INDEX
        sheet.

        Like Nippon, MO ships ONE workbook for all schemes, so every entry
        in the returned dict maps to the same URL.
        """
        url = self._select_consolidated_url(ym)
        if url is None:
            log.warning("holdings.motilal_oswal.no_url_for_ym", ym=ym)
            return {}

        path = self._cached_path(ym)
        if not path.exists():
            download_to(url, path)

        out: dict[str, str] = {}
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 — surface payload issue
            log.error(
                "holdings.motilal_oswal.workbook_open_failed",
                ym=ym, path=str(path), err=str(e),
            )
            return {}
        try:
            if "INDEX" not in wb.sheetnames:
                log.warning(
                    "holdings.motilal_oswal.no_index_sheet",
                    ym=ym, sheetnames=wb.sheetnames[:5],
                )
                return {}
            idx_ws = wb["INDEX"]
            sheet_set = set(wb.sheetnames)
            for row in idx_ws.iter_rows(values_only=True):
                cells = list(row) + [None] * max(0, 5 - len(row))
                raw_name = cells[3]   # col D: Fund Name
                code = cells[4]       # col E: Fund Code
                # Resolve the code to a real sheet (tolerating the Y0/YO
                # digit-zero typo); skip rows that name no sheet (header etc.)
                if _resolve_sheet_code(code, sheet_set) is None:
                    continue
                name = _extract_scheme_name(raw_name)
                if not name:
                    continue
                out.setdefault(_emit_name(name), url)
        finally:
            wb.close()

        log.info(
            "holdings.motilal_oswal.discover",
            ym=ym, n_schemes=len(out), url=url,
        )
        return out

    def _select_consolidated_url(self, ym: str) -> str | None:
        """Find the monthly consolidated workbook's absolute URL for data
        month ``ym``.

        Search the publish-month folder (data + 1) first — the file's normal
        home — then fall back to the data-month folder. In both, keep only
        the result whose TITLE is the canonical monthly portfolio for ``ym``.
        """
        candidates: list[tuple[int, str]] = []
        py, pm_short, _ = _ym_parts(_publish_ym(ym))
        dy, dm_short, _ = _ym_parts(ym)
        # (year, short-month) folders to scan, in priority order.
        for year, mon in [(py, pm_short), (dy, dm_short)]:
            for r in _query_folder(year, mon):
                title = r.get("title") or ""
                rel = r.get("path") or ""
                if not rel.lower().endswith((".xls", ".xlsx")):
                    continue
                if _is_monthly_title(title, ym):
                    candidates.append((year, rel))
            if candidates:
                break
        if not candidates:
            return None
        rel = candidates[0][1]
        # Absolutize + percent-encode the DAM path (it contains spaces).
        from urllib.parse import quote

        if rel.startswith("http"):
            return rel
        return _HOST + quote(rel)

    def _cached_path(self, ym: str) -> Path:
        return paths.holdings_excel_raw(
            self.amc_slug, ym, _consolidated_filename(ym)
        )

    # ------------------------------------------------------------------
    # Download
    # ------------------------------------------------------------------

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Return the cached consolidated workbook for ``ym``.

        MO's workbook is shared across all schemes, so ``scheme_filename``
        is intentionally ignored — every scheme resolves to the same file.
        """
        path = self._cached_path(ym)
        if path.exists():
            return path
        return download_to(url, path)

    # ------------------------------------------------------------------
    # Parse
    # ------------------------------------------------------------------

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        """Walk the scheme's sheet inside the consolidated workbook and
        yield ISIN-bearing holding rows.

        We look up the scheme's ``YOnn`` sheet code from the INDEX sheet,
        then run the shared generic SEBI parser on that one sheet (it
        auto-detects the header, columns, and percent unit).
        """
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            sheet_code = self._find_sheet_for_scheme(wb, scheme_name_printed)
            if sheet_code is None:
                log.warning(
                    "holdings.motilal_oswal.scheme_not_in_index",
                    scheme=scheme_name_printed, ym=ym,
                )
                return
            ws = wb[sheet_code]
            recs = _parse_one_sheet(ws)
            if not recs:
                log.warning(
                    "holdings.motilal_oswal.empty_sheet",
                    scheme=scheme_name_printed, sheet=sheet_code, ym=ym,
                )
                return
            for r in recs:
                yield ParsedHoldingRecord(
                    scheme_name_printed=scheme_name_printed,
                    security_name=r.security_name,
                    weight_pct=r.weight_pct,
                    isin=r.isin,
                    instrument_type=r.instrument_type,
                    source_amc=self.amc_slug,
                )
        finally:
            wb.close()

    @staticmethod
    def _find_sheet_for_scheme(
        wb: openpyxl.Workbook, scheme_name_printed: str
    ) -> str | None:
        """Reverse-lookup the ``YOnn`` sheet code from the INDEX sheet.

        Match is case-insensitive on the cleaned scheme name. Returns the
        sheet-code string, or None if not found.
        """
        if "INDEX" not in wb.sheetnames:
            return None
        # Reverse the discovery-time disambiguation suffix so the cleaned
        # INDEX name compares equal again.
        lookup = _strip_disambiguation(scheme_name_printed.strip())
        target = re.sub(r"\s+", " ", lookup.lower())
        idx_ws = wb["INDEX"]
        sheet_set = set(wb.sheetnames)
        for row in idx_ws.iter_rows(values_only=True):
            cells = list(row) + [None] * max(0, 5 - len(row))
            raw_name = cells[3]
            code = cells[4]
            sheet_code = _resolve_sheet_code(code, sheet_set)
            if sheet_code is None:
                continue
            name = _extract_scheme_name(raw_name)
            if name and name.lower() == target:
                return sheet_code
        return None
