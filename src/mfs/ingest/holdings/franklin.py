"""Franklin Templeton (India) monthly portfolio holdings adapter (Phase 5).

Like Nippon — and UNLIKE HDFC/SBI (one Excel per scheme) — Franklin
Templeton India publishes ONE consolidated workbook per month covering every
FTMF scheme, with one worksheet per scheme.

Discovery (reverse-engineered from the Angular SPA at
``franklintempletonindia.com``):

The Reports/Disclosures page is a JS SPA whose download catalog is served by
a single JSON endpoint::

    GET /api/literature/v1/responseLitJson?type=report

The response is ``{"FirstDropDown":[ {id, dataRecords:{linkdata:[...]}}, ...]}``.
The monthly-portfolio category is the FirstDropDown entry with
``id == "MONTHLY-PORTFOLIO-DSCLR"``; its ``dataRecords.linkdata`` is the list
of monthly workbooks, one per month, each a record like::

    {"dctermsTitle": "ISIN as on 30 April 2026",
     "frkReferenceDate": "2026-04-30",
     "literatureHref": "/en-in/monthly-portfolio-dsclr/<uuid>/"
                       "Monthly-Portfolio-ISIN-30-Apr-2026.xlsx"}

``frkReferenceDate`` is the data-month-end date (NOT the publish date), so we
pick the record whose ``frkReferenceDate`` falls inside the requested data
month ``ym``. The bare ``literatureHref`` is a download path that must be
prefixed with ``/download`` (mirrors the SPA's own ``"download"+href`` link
builder), giving the absolute URL::

    https://www.franklintempletonindia.com/download/en-in/monthly-portfolio-
        dsclr/<uuid>/Monthly-Portfolio-ISIN-30-Apr-2026.xlsx

The ``<uuid>`` and exact filename change every month, so we never hardcode
them — we always resolve them from the live JSON, keyed by ``ym``.

Workbook structure (validated against the 30-Apr-2026 file, ~1.2 MB):
- One worksheet per scheme; sheet name is the internal scheme code
  (e.g. ``FIEF`` = Flexi Cap, ``FIFEF`` = Focused Equity, ``FILCF`` =
  Large Cap). There is NO "Index" sheet — the scheme name lives in row 0.
- Per-sheet layout:
    row 0: scheme-name banner in col A (e.g. "Franklin India Flexi Cap Fund
           ( Formerly known as Franklin India Equity Fund)").
    row 2: "Portfolio Statement as on April 30, 2026".
    row 3: header — col A="ISIN Number", col B="Name of the Instrument",
           col C="Industry Classification / Rating", col D="Quantity",
           col E="Market Value ... (Rs. in Lakhs)", col F="% to Net Assets".
    row 4+: section banners ("Equity & Equity related", "(a) Listed /
            awaiting listing on Stock Exchanges", "Debt Instruments",
            "Money Market Instruments", ...) interleaved with holding rows.
- ``% to Net Assets`` is already a PERCENT (8.53 == 8.53%), summing to ~100.

This layout matches the generic SEBI shape exactly, so we reuse
``parse_sebi_excel``'s single-sheet parser (``_parse_one_sheet``) rather than
re-implementing column detection. We only add the consolidated-workbook
plumbing: map each scheme name to its sheet code, and at parse time select the
one sheet whose row-0 name matches the requested ``scheme_name_printed``.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path

import openpyxl

from mfs import paths
from mfs.ingest.holdings._base import HoldingsAdapter
from mfs.ingest.holdings._generic import _parse_one_sheet
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to, fetch_bytes
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_HOST = "https://www.franklintempletonindia.com"
# Single JSON endpoint backing the Reports/Disclosures download SPA. The
# monthly-portfolio catalog lives under the "report" type (not "download").
_LISTING_API = (
    "https://www.franklintempletonindia.com/api/literature/v1/responseLitJson"
    "?type=report"
)
# Bare literatureHref values are download paths; the SPA prefixes them with
# "download" to form the real file URL.
_DOWNLOAD_PREFIX = "https://www.franklintempletonindia.com/download"
# FirstDropDown entry id carrying the monthly portfolio workbooks.
_MONTHLY_CATEGORY_ID = "MONTHLY-PORTFOLIO-DSCLR"


def _href_to_url(href: str) -> str:
    """Turn a relative literatureHref into an absolute download URL.

    Handles the rare already-absolute case defensively; otherwise prefixes
    the bare ``/en-in/...`` path with the ``/download`` segment.
    """
    if href.startswith("http"):
        return href
    if not href.startswith("/"):
        href = "/" + href
    return _DOWNLOAD_PREFIX + href


def _master_excel_filename(ym: str) -> str:
    """Canonical local-cache filename for the consolidated monthly Excel."""
    return f"FTMF-MONTHLY-PORTFOLIO-{ym}.xlsx"


def _clean_scheme_name(raw: object) -> str | None:
    """Normalize a row-0 scheme-name banner to a stable printed name.

    Collapses whitespace and strips Franklin's trailing footnote markers
    ("^", "$$~~", trailing "*"). The "(Formerly known as ...)" parenthetical
    is intentionally KEPT — it disambiguates renamed schemes and the
    downstream fuzzy matcher handles it.
    """
    if not isinstance(raw, str):
        return None
    s = re.sub(r"\s+", " ", raw).strip()
    # Strip trailing presentation markers Franklin appends to scheme names.
    s = re.sub(r"[\^*~$]+$", "", s).strip()
    return s or None


def _data_month_matches(ref_date: str, ym: str) -> bool:
    """True if frkReferenceDate (YYYY-MM-DD) falls inside data month ym."""
    if not isinstance(ref_date, str) or len(ref_date) < 7:
        return False
    return ref_date[:7] == ym


def _select_master_record(listing: dict, ym: str) -> dict | None:
    """Find the monthly-portfolio linkdata record for data month ``ym``.

    Returns the record dict (carrying ``literatureHref`` + ``frkReferenceDate``)
    or None if no workbook is published for that month yet.
    """
    for cat in listing.get("FirstDropDown", []):
        if cat.get("id") != _MONTHLY_CATEGORY_ID:
            continue
        records = (cat.get("dataRecords") or {}).get("linkdata") or []
        for rec in records:
            href = rec.get("literatureHref") or ""
            if ".xls" not in href.lower():
                continue
            if _data_month_matches(str(rec.get("frkReferenceDate", "")), ym):
                return rec
        return None  # category found but no record for ym
    return None


@register_adapter
class FranklinHoldingsAdapter(HoldingsAdapter):
    amc_slug = "franklin"
    source_label = "Franklin Templeton Mutual Fund"

    # ------------------------------------------------------------------
    # Discovery
    # ------------------------------------------------------------------

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Resolve the consolidated workbook for ``ym``, cache it, and
        enumerate the per-scheme sheets from their row-0 names.

        Because Franklin ships a SINGLE workbook for all schemes, every
        returned entry maps to the same URL; ``parse_excel`` later selects
        the matching sheet by scheme name.
        """
        import json

        raw = fetch_bytes(_LISTING_API)
        try:
            listing = json.loads(raw.decode("utf-8", errors="replace"))
        except json.JSONDecodeError as e:
            log.error("holdings.franklin.listing_parse_failed", ym=ym, err=str(e))
            return {}

        rec = _select_master_record(listing, ym)
        if rec is None:
            log.warning("holdings.franklin.no_url_for_ym", ym=ym)
            return {}
        master_url = _href_to_url(rec["literatureHref"])

        master_path = self._cached_master_path(ym)
        if not master_path.exists():
            download_to(master_url, master_path)

        out: dict[str, str] = {}
        try:
            wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 — surface any payload issue
            log.error(
                "holdings.franklin.workbook_open_failed",
                ym=ym, path=str(master_path), err=str(e),
            )
            return {}
        try:
            for sheet in wb.worksheets:
                name = self._sheet_scheme_name(sheet)
                if name:
                    out.setdefault(name, master_url)
        finally:
            wb.close()

        log.info(
            "holdings.franklin.discover",
            ym=ym, n_schemes=len(out), master_url=master_url,
        )
        return out

    # ------------------------------------------------------------------
    # Download
    # ------------------------------------------------------------------

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Return the cached consolidated workbook for this month.

        ``scheme_filename`` is intentionally ignored — all schemes share the
        one workbook. Download is idempotent in case fetch is called without
        going through discover first.
        """
        master_path = self._cached_master_path(ym)
        if master_path.exists():
            return master_path
        return download_to(url, master_path)

    # ------------------------------------------------------------------
    # Parse
    # ------------------------------------------------------------------

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        """Select the worksheet whose row-0 name matches ``scheme_name_printed``
        and yield its ISIN-bearing holdings via the shared generic parser."""
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            target = scheme_name_printed.strip().lower()
            ws = None
            for sheet in wb.worksheets:
                name = self._sheet_scheme_name(sheet)
                if name and name.lower() == target:
                    ws = sheet
                    break
            if ws is None:
                log.warning(
                    "holdings.franklin.scheme_sheet_not_found",
                    scheme=scheme_name_printed, ym=ym,
                )
                return
            recs = _parse_one_sheet(ws)
            if not recs:
                return
            for r in recs:
                yield ParsedHoldingRecord(
                    scheme_name_printed=scheme_name_printed,
                    security_name=r.security_name,
                    weight_pct=r.weight_pct,
                    isin=r.isin,
                    instrument_type=r.instrument_type,
                    source_amc=self.amc_slug,
                )
        finally:
            wb.close()

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _sheet_scheme_name(ws):  # type: ignore[no-untyped-def]
        """Pull the printed scheme name from a sheet's row-0 banner."""
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            if row:
                return _clean_scheme_name(row[0])
        return None

    def _cached_master_path(self, ym: str) -> Path:
        return paths.holdings_excel_raw(
            self.amc_slug, ym, _master_excel_filename(ym)
        )
