"""Bajaj Finserv Mutual Fund monthly portfolio holdings adapter (Phase 5).

Like Nippon (and unlike HDFC's one-Excel-per-scheme), Bajaj Finserv publishes
ONE consolidated ``.xlsx`` workbook per month covering every scheme — one sheet
per scheme, sheet name == the fund's internal code (e.g. ``BFFLX`` = Flexi Cap),
the printed scheme name in row 0 col B.

Discovery (WordPress + admin-ajax SPA)
--------------------------------------
The downloads page (https://www.bajajamc.com/downloads) is a WordPress site
whose "Portfolio > Monthly Portfolio" accordion is rendered by a custom
``bajaj-downloads`` plugin. The links are NOT in the static HTML; they come
from an admin-ajax POST:

    POST https://www.bajajamc.com/wp-admin/admin-ajax.php
    form body:
        action     = bajaj_get_downloads
        nonce      = <wp nonce read off `var bajajDownloads = {...}` in the page>
        section_id = 757            (the "Monthly Portfolio" accordion section)
        year       = <FY label>     (Indian FINANCIAL year, e.g. '2026-27')
        month      = <Month name>   (e.g. 'April')
    -> {"success":true,"data":{"html":"<div class=bd-download-row>…<a href=…
        .xlsx>…","count":N}}

CRITICAL: the year dropdown is keyed by the Indian FINANCIAL year (Apr-Mar),
NOT the calendar year. April 2026 data lives under FY ``2026-27``; a
Jan/Feb/Mar data month lives under the PRIOR FY (e.g. Feb 2026 -> ``2025-26``).
``_fy_label`` maps a data ym to the right FY tab.

The WP nonce is issued to anonymous visitors (no login) but is bound to the
AWSALB session cookie, so we drive the whole flow inside one ``httpx.Client``:
GET the page to seed the cookie + read the nonce, then POST the listing action.

The card title / filename encodes the DATA-month-end date
("…_Monthly Portfolio as on 30 Apr 2026.xlsx") — we filter the returned rows to
the row whose date falls in the requested data month, so the FY/month inputs
plus that filter make next month work unchanged.

Workbook layout (validated against the April 2026 file, ~2.1 MB, 24 sheets)
--------------------------------------------------------------------------
- No "Index" sheet; sheets 0..N are one-per-scheme, named by internal code.
- Per-scheme sheet:
    row 0: col A = code, col B = printed scheme name
           ("Bajaj Finserv Flexi Cap Fund").
    row 2: "Monthly Portfolio Statement as on April 30…".
    row 3: header — col B='Name of the Instrument', col C='ISIN',
           col D='Industry', col E='Quantity',
           col F='Market/Fair Value (Rs. in Lakhs)', col G='% to Net Assets',
           col H='YTM~', col I='YTC^'.
    row 4+: section banners ('Equity & Equity related', '(a) Listed / awaiting
           listing on Stock Exchanges') interleaved with ISIN-bearing rows.
- CRITICAL UNIT NOTE: '% to Net Assets' is stored as a **fraction** (0.0499 ==
  4.99%), like Nippon/Axis/Baroda BNP.

Because each sheet is the canonical SEBI layout, we reuse the shared
``parse_sebi_excel`` — it auto-detects the header row, the ISIN/name/weight
columns, AND the fraction-vs-percent weight unit. We only need to point it at
the correct sheet (looked up by the row-0 scheme name).
"""

from __future__ import annotations

import html as _html
import json
import re
from collections.abc import Iterable
from pathlib import Path

import httpx
import openpyxl

from mfs import paths
from mfs.config import get_settings
from mfs.ingest.holdings._generic import GenericHoldingsAdapter, parse_sebi_excel
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_ORIGIN = "https://www.bajajamc.com"
_DOWNLOADS_PAGE = f"{_ORIGIN}/downloads"
_AJAX_URL = f"{_ORIGIN}/wp-admin/admin-ajax.php"

# The "Monthly Portfolio" accordion's section id (read off the page's
# data-section-id; the accordion carries data-filter="year_month"). Stable.
_SECTION_ID = "757"

_MONTHS = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)

# Short month tokens as they appear in the card title / filename
# ("…as on 30 Apr 2026.xlsx").
_MONTH_ABBR = (
    "Jan", "Feb", "Mar", "Apr", "May", "Jun",
    "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
)

# `var bajajDownloads = {"ajaxUrl":"…","nonce":"<hex>"};`
_NONCE_RE = re.compile(
    r'bajajDownloads\s*=\s*\{[^}]*"nonce"\s*:\s*"(?P<nonce>[0-9a-f]+)"',
    re.IGNORECASE,
)

# A download row in the AJAX html fragment: the .xlsx href plus the "as on
# DD Mon YYYY" date in the title. The href is the only thing we need; the date
# is parsed to confirm it matches the requested data month.
_HREF_RE = re.compile(
    r'href="(?P<url>https://media\.bajajamc\.com/[^"]+?\.xlsx)"',
    re.IGNORECASE,
)
_AS_ON_RE = re.compile(
    r'as on\s+(?P<day>\d{1,2})\s+(?P<mon>[A-Za-z]{3,9})\s+(?P<year>\d{4})',
    re.IGNORECASE,
)

_UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"
)


def _fy_label(ym: str) -> str:
    """data ym -> Indian financial-year dropdown label (Apr-Mar).

    April..December -> 'YYYY-(YY+1)'; January..March -> '(YYYY-1)-YY'.
    e.g. '2026-04' -> '2026-27'; '2026-02' -> '2025-26'.
    """
    y, m = map(int, ym.split("-"))
    start = y if m >= 4 else y - 1
    return f"{start:04d}-{(start + 1) % 100:02d}"


def _data_month_name(ym: str) -> str:
    """data ym -> full English month name (the `month` dropdown value)."""
    _, m = map(int, ym.split("-"))
    return _MONTHS[m - 1]


def _matches_data_month(fragment_around_href: str, ym: str) -> bool:
    """True if the row's 'as on DD Mon YYYY' date falls in data month ``ym``.

    Defensive against the AJAX listing ever returning more than the single
    consolidated workbook for a month (it currently returns exactly one).
    """
    m = _AS_ON_RE.search(fragment_around_href)
    if not m:
        return False
    y, mo = map(int, ym.split("-"))
    mon = m.group("mon").lower()
    want_full = _MONTHS[mo - 1].lower()
    want_abbr = _MONTH_ABBR[mo - 1].lower()
    return int(m.group("year")) == y and mon in (want_full, want_abbr)


def _master_filename(ym: str) -> str:
    """Canonical local-cache name for the month's consolidated workbook."""
    return f"Bajaj-Finserv-Mutual-Fund_Monthly-Portfolio-{ym}.xlsx"


@register_adapter
class BajajFinservHoldingsAdapter(GenericHoldingsAdapter):
    amc_slug = "bajaj_finserv"
    source_label = "Bajaj Finserv Mutual Fund"

    # ------------------------------------------------------------------
    # Discovery
    # ------------------------------------------------------------------

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Find the month's consolidated workbook via admin-ajax, cache it,
        and enumerate the per-scheme sheets.

        Bajaj publishes ONE workbook for all schemes, so every entry maps to
        the same URL (Nippon-style); ``parse_excel`` later picks the right
        sheet by the printed scheme name.
        """
        master_url = self._discover_master_url(ym)
        if master_url is None:
            log.warning("holdings.bajaj_finserv.no_url_for_ym", ym=ym)
            return {}

        master_path = self._cached_master_path(ym)
        if not master_path.exists():
            download_to(master_url, master_path)

        out: dict[str, str] = {}
        try:
            wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 — surface any payload issue
            log.error(
                "holdings.bajaj_finserv.workbook_open_failed",
                ym=ym, path=str(master_path), err=str(e),
            )
            return {}
        try:
            for ws in wb.worksheets:
                name = self._sheet_scheme_name(ws)
                if name:
                    out.setdefault(name, master_url)
        finally:
            wb.close()

        log.info(
            "holdings.bajaj_finserv.discover",
            ym=ym, n_schemes=len(out), master_url=master_url,
        )
        return out

    def _discover_master_url(self, ym: str) -> str | None:
        """POST the admin-ajax listing for the data month's FY+month tab and
        return the consolidated workbook URL whose 'as on' date is in ``ym``."""
        fy = _fy_label(ym)
        month = _data_month_name(ym)
        s = get_settings()
        with httpx.Client(
            timeout=s.http_timeout,
            follow_redirects=True,
            headers={"User-Agent": _UA},
        ) as c:
            page = c.get(_DOWNLOADS_PAGE)
            page.raise_for_status()
            nm = _NONCE_RE.search(page.text)
            if not nm:
                log.error("holdings.bajaj_finserv.no_nonce", ym=ym)
                return None
            nonce = nm.group("nonce")

            r = c.post(
                _AJAX_URL,
                data={
                    "action": "bajaj_get_downloads",
                    "nonce": nonce,
                    "section_id": _SECTION_ID,
                    "year": fy,
                    "month": month,
                },
                headers={
                    "X-Requested-With": "XMLHttpRequest",
                    "Referer": _DOWNLOADS_PAGE,
                    "Accept": "application/json, text/javascript, */*; q=0.01",
                },
            )
            r.raise_for_status()
            try:
                payload = r.json()
            except json.JSONDecodeError:
                log.error(
                    "holdings.bajaj_finserv.bad_json",
                    ym=ym, head=r.text[:200],
                )
                return None

        if not payload.get("success"):
            log.warning("holdings.bajaj_finserv.ajax_unsuccessful", ym=ym, fy=fy)
            return None
        fragment = _html.unescape((payload.get("data") or {}).get("html") or "")

        # Walk every .xlsx href and keep the first whose row's "as on" date
        # matches the requested data month. Each row's date sits just before
        # its href in the fragment.
        for m in _HREF_RE.finditer(fragment):
            row_text = fragment[max(0, m.start() - 300): m.end()]
            if _matches_data_month(row_text, ym):
                return m.group("url")
        log.warning(
            "holdings.bajaj_finserv.no_matching_row",
            ym=ym, fy=fy, month=month, head=fragment[:200],
        )
        return None

    # ------------------------------------------------------------------
    # Fetch
    # ------------------------------------------------------------------

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Return the cached consolidated workbook for ``ym``.

        Bajaj's master Excel is shared across all schemes, so
        ``scheme_filename`` is ignored. We download here (idempotent) in case
        ``_run.py`` invoked us without going through discover first.
        """
        master_path = self._cached_master_path(ym)
        if master_path.exists():
            return master_path
        return download_to(url, master_path)

    def _cached_master_path(self, ym: str) -> Path:
        return paths.holdings_excel_raw(self.amc_slug, ym, _master_filename(ym))

    # ------------------------------------------------------------------
    # Parse
    # ------------------------------------------------------------------

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        """Locate the scheme's sheet (by its row-0 printed name) inside the
        consolidated workbook and delegate to the shared ``parse_sebi_excel``.

        ``parse_sebi_excel`` auto-detects the header row, the ISIN/name/weight
        columns, and the fraction-vs-percent weight unit, so no bespoke
        column logic is needed — we only resolve which sheet to read.
        """
        sheet_index = self._find_sheet_index(excel_path, scheme_name_printed)
        if sheet_index is None:
            log.warning(
                "holdings.bajaj_finserv.scheme_not_found",
                scheme=scheme_name_printed, ym=ym,
            )
            return
        yield from parse_sebi_excel(
            excel_path, scheme_name_printed, self.amc_slug,
            sheet_index=sheet_index, expect_ym=ym,
        )

    @staticmethod
    def _sheet_scheme_name(ws) -> str | None:
        """Pull the printed scheme name from a sheet's row 0 col B."""
        try:
            first = next(ws.iter_rows(values_only=True))
        except StopIteration:
            return None
        name = first[1] if len(first) > 1 else None
        if not isinstance(name, str):
            return None
        name = re.sub(r"\s+", " ", name).strip()
        return name or None

    def _find_sheet_index(
        self, excel_path: Path, scheme_name_printed: str,
    ) -> int | None:
        """Return the 0-based sheet index whose row-0 name matches (case-
        insensitive), or None."""
        target = re.sub(r"\s+", " ", scheme_name_printed.strip()).lower()
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            for i, ws in enumerate(wb.worksheets):
                name = self._sheet_scheme_name(ws)
                if name and name.lower() == target:
                    return i
        finally:
            wb.close()
        return None
