"""Old Bridge Mutual Fund monthly portfolio holdings adapter (Phase 5).

Old Bridge is a small AMC (three schemes: Flexi Cap, Focused, Arbitrage). It
publishes ONE SEBI-format Excel per scheme per month on a STATIC HTML page —
the statutory-disclosures page server-renders every download anchor as a
plain ``/uploads/...xlsx`` link (no SPA / JSON API):

    https://oldbridgemf.com/statutory-disclosures.html

The page lists, alongside the per-scheme Monthly Portfolios, several other
.xlsx document families we must NOT ingest: Half-Yearly Portfolios
(``HY_*`` / ``Half_Yearly_*``), Half-Yearly Financials, Scheme Dashboards,
and AMFI AAUM workbooks (``*Average_Asset*``). The disclosures-page filenames
are NOT stable month-over-month — across the archive the same fund's monthly
portfolio has appeared as ``Old_Bridge_Flexi_Cap_Fund_Apr_26_Portfolio_*``,
``Flexi_Cap_Monthly_Portfolio_March_26_*``, ``Monthly_Portfolio_FE_Dec2025_*``,
``Monthly_Portfolio_October_25_*`` (a combined-looking name), etc. — and a
content-hash suffix is appended to every upload.

Because the filename scheme tokens are unreliable, discovery does NOT try to
read the fund name out of the filename. Instead it:
  1. scrapes every ``.xlsx`` link on the page,
  2. drops the non-monthly-portfolio families by keyword,
  3. keeps files whose filename carries the requested DATA month (month name
     or 3-letter abbrev + 2- or 4-digit year),
  4. OPENS each surviving workbook and reads the scheme name from the
     in-sheet banner (row 2/3, e.g. "Old Bridge Flexi Cap Fund"), which is
     authoritative and stable.

The in-sheet name is then fuzzy-matched to scheme_master by the orchestrator.
A file that carries no recognizable fund banner (e.g. a combined workbook, or
a month where Old Bridge shipped one file for all schemes) is skipped rather
than guessed — we never fabricate a scheme assignment.

Per-scheme Excel layout (validated against Flexi Cap & Focused, April 2026):
  - One sheet, named after the internal scheme code ('OBFLX' / 'OBFCE').
  - row 0: '<CODE> | Old Bridge Mutual Fund' banner.
  - row 1: 'Monthly Portfolio Statement as on ...'.
  - row 2: the printed scheme name ('Old Bridge Flexi Cap Fund').
  - row 4: header — col B='Name of the Instrument', col C='ISIN',
    col D='Industry*', col E='Quantity', col F='Market/Fair Value (Rs. in
    Lakhs)', col G='% to Net Assets', col H='YTM'.
  - row 5+: section banners ('Equity & Equity related', '(a) Listed /
    awaiting listing on Stock Exchanges', ...) interleaved with holding rows.

The header is a textbook SEBI layout, and '% to Net Assets' is stored as a
PERCENT (Gujarat Ambuja at 3.34 == 3.34%), so the shared ``parse_sebi_excel``
handles both parsing and weight-unit detection. We override nothing on the
parse side — only discovery is bespoke.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from mfs import paths
from mfs.ingest.holdings._generic import GenericHoldingsAdapter
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to, fetch_bytes
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_DISCLOSURE_PAGE = "https://oldbridgemf.com/statutory-disclosures.html"
_BASE = "https://oldbridgemf.com"

# Any .xlsx URL on the page, absolute or root-relative.
_XLSX_RE = re.compile(
    r'(?P<url>(?:https://oldbridgemf\.com)?/[^"\'\s)]+?\.xlsx?)',
    re.IGNORECASE,
)

# Document families on the page that are NOT per-scheme monthly portfolios.
# Matched against the lowercased filename. Order/specificity doesn't matter;
# any hit excludes the file.
_EXCLUDE_KEYWORDS = (
    "half_yearly",
    "half yearly",
    "hy_",          # HY_OBMF_Portfolio_*, HY_Portfolio_*
    "financial",
    "average_asset",
    "average_assets",
    "assets_under_management",
    "asset_under_management",
    "aaum",
    "scheme_dashboard",
    "dashboard",
    "scheme_performance",
    "performance",
    "addend",       # addenda
    "factsheet",
    "fact_sheet",
)

_MONTHS = [
    "january", "february", "march", "april", "may", "june",
    "july", "august", "september", "october", "november", "december",
]
_MONTH_ABBR = [m[:3] for m in _MONTHS]  # jan, feb, mar, ...


def _month_tokens(ym: str) -> list[re.Pattern]:
    """Build regexes that match the DATA month inside a filename.

    Old Bridge's filenames spell the month as either the full name or the
    3-letter abbreviation, and the year as either 2 or 4 digits, with the two
    separated by an underscore or directly adjacent (``Apr_26``, ``April_2026``,
    ``Apr2026``, ``April26``, ``Dec2025``). We require the month token and the
    year token to be adjacent (optionally separated by ``_``) so that, e.g.,
    'March_25' in a 2025 file never matches a request for May.
    """
    y, m = map(int, ym.split("-"))
    full = _MONTHS[m - 1]
    abbr = _MONTH_ABBR[m - 1]
    yy = f"{y % 100:02d}"
    yyyy = f"{y:04d}"
    # full|abbr , then optional '_' , then 4-digit OR 2-digit year as a token.
    return [
        re.compile(
            rf"(?:^|[_\- ])(?:{full}|{abbr})_?(?:{yyyy}|{yy})(?:[_\-. ]|$)",
            re.IGNORECASE,
        )
    ]


def _is_excluded(filename_lower: str) -> bool:
    return any(k in filename_lower for k in _EXCLUDE_KEYWORDS)


def _scheme_name_from_workbook(path: Path) -> str | None:
    """Read the printed scheme name from the in-sheet banner.

    The banner ('Old Bridge <Fund> Fund') sits in the first few rows above the
    'Name of the Instrument' header. We scan the leading rows for a string that
    starts with 'Old Bridge' and is not the AMC name / the 'Monthly Portfolio
    Statement ...' line. Returns None if no recognizable banner is found
    (caller skips the file rather than guessing).
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — corrupt / unexpected payload
        log.warning("holdings.old_bridge.workbook_open_failed",
                    path=str(path), err=str(e))
        return None
    try:
        ws = wb[wb.sheetnames[0]]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 8:  # banner is always in the preamble, well above holdings
                break
            for cell in row:
                if not isinstance(cell, str):
                    continue
                s = re.sub(r"\s+", " ", cell).strip()
                low = s.lower()
                if not low.startswith("old bridge"):
                    continue
                if "mutual fund" in low:  # AMC name, not a scheme
                    continue
                if "portfolio statement" in low or "monthly portfolio" in low:
                    continue
                if "fund" in low:  # scheme banners all contain 'Fund'
                    return s
    finally:
        wb.close()
    return None


@register_adapter
class OldBridgeHoldingsAdapter(GenericHoldingsAdapter):
    amc_slug = "old_bridge"
    source_label = "Old Bridge Mutual Fund"

    # parse_excel / fetch_excel inherited from GenericHoldingsAdapter:
    # the shared parse_sebi_excel auto-detects the header, ISIN/name/weight
    # columns, and the percent weight unit on sheet 0.

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Return {printed_scheme_name: absolute_excel_url} for data month ``ym``.

        Scrapes the static disclosures page, excludes the non-monthly-portfolio
        document families, keeps files whose filename carries ``ym``'s month,
        and reads each surviving workbook's in-sheet banner for the
        authoritative scheme name.
        """
        html = fetch_bytes(_DISCLOSURE_PAGE).decode("utf-8", errors="replace")
        month_pats = _month_tokens(ym)

        seen_urls: set[str] = set()
        candidate_urls: list[str] = []
        for m in _XLSX_RE.finditer(html):
            url = m.group("url")
            if url.startswith("/"):
                url = _BASE + url
            if url in seen_urls:
                continue
            seen_urls.add(url)
            filename = url.rsplit("/", 1)[-1]
            low = filename.lower()
            if _is_excluded(low):
                continue
            if not any(p.search(low) for p in month_pats):
                continue
            candidate_urls.append(url)

        out: dict[str, str] = {}
        for url in candidate_urls:
            filename = url.rsplit("/", 1)[-1]
            try:
                local = self._fetch_cached(url, filename, ym)
            except Exception as e:  # noqa: BLE001 — skip a flaky download
                log.warning("holdings.old_bridge.download_failed",
                            url=url, ym=ym, err=str(e))
                continue
            scheme = _scheme_name_from_workbook(local)
            if not scheme:
                log.warning("holdings.old_bridge.no_scheme_banner",
                            url=url, ym=ym)
                continue
            # First file for a given scheme name wins (page lists newest first).
            out.setdefault(scheme, url)

        log.info("holdings.old_bridge.discover",
                 ym=ym, n_candidates=len(candidate_urls), n_schemes=len(out))
        return out

    def _fetch_cached(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Download (cached) under the raw-holdings cache. Reused by both
        discovery (to read the banner) and the orchestrator's fetch step."""
        out = paths.holdings_excel_raw(self.amc_slug, ym, scheme_filename)
        if out.exists():
            return out
        return download_to(url, out)

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        # The orchestrator passes scheme_filename="<printed name>.xlsx"; cache
        # under the source filename instead so discovery's already-downloaded
        # copy is reused (and so two schemes never collide on a cache path).
        source_filename = url.rsplit("/", 1)[-1]
        return self._fetch_cached(url, source_filename, ym)
