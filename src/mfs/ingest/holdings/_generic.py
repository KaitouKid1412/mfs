"""Generic SEBI monthly-portfolio Excel parser + config-driven adapter.

SEBI mandates that every AMC publish a monthly portfolio statement in a
broadly standardized layout: a header row carrying columns that include
``ISIN``, a ``Name of the Instrument`` column, an industry/rating column,
Quantity, Market Value, and a ``% to NAV`` weight column (AMCs variously
print ``% to AUM`` / ``% of AUM`` / ``% to Net Assets`` / ``% of Net
Assets``).

The bespoke ``hdfc`` / ``sbi`` / ``nippon`` adapters hard-code column
offsets and the weight unit (percent vs fraction) because they predate this
module. Everything written from Phase 5 onward should reuse
``parse_sebi_excel`` instead: it DETECTS the header row, the ISIN / name /
weight column positions, and the weight unit at parse time, so one parser
handles the long tail of AMCs whose layouts differ only cosmetically.

A scheme's complete portfolio sums to ~100 (percent layout) or ~1.0
(fraction layout); we sum the raw weights and scale fractions up by 100.
This is more robust than a per-row magnitude heuristic (a fund whose top
holding is 1.2% would be misread as a fraction by a max-based rule).
"""

from __future__ import annotations

import datetime as _dt
import re
from collections.abc import Callable, Iterable
from pathlib import Path
from urllib.parse import unquote, urljoin

import openpyxl

from mfs import paths
from mfs.errors import StatementDateMismatchError
from mfs.ingest.holdings._base import HoldingsAdapter
from mfs.io.http import download_to, fetch_bytes
from mfs.schemas import ParsedHoldingRecord

_ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}\d$")


def is_isin(s: str) -> bool:
    """Strict ISIN check: 2 letters + 9 alphanumeric + 1 check digit."""
    return bool(_ISIN_RE.match(s))


def _norm(cell: object) -> str:
    """Lowercase, collapse-whitespace, strip-punctuation view of a cell for
    header matching. Returns '' for non-text / empty cells."""
    if not isinstance(cell, str):
        return ""
    s = cell.replace("\n", " ").strip().lower()
    s = re.sub(r"[^a-z0-9%]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


# Header-cell predicates. We match on normalized text so '% to NAV',
# '% to AUM', '% of AUM', '% to Net Assets' all resolve to the weight column.
def _is_isin_header(s: str) -> bool:
    return s == "isin" or s.startswith("isin ") or s.endswith(" isin")


def _is_weight_header(s: str) -> bool:
    if "%" not in s:
        return False
    return any(
        k in s
        for k in ("to nav", "of nav", "to aum", "of aum",
                  "to net asset", "of net asset", "to total", "of total")
    )


def _is_name_header(s: str) -> bool:
    if any(
        k in s
        for k in ("name of the instrument", "name of instrument",
                  "instrument issuer", "name of the issuer", "security name",
                  "company name", "name of instrument issuer")
    ) or s in ("instrument", "name"):
        return True
    # Broader catch for the "Company/Issuer/Instrument Name" header variant
    # (ICICI/Mirae ETFs): any non-weight header naming the instrument/issuer.
    # Excludes the weight column ('%') and the industry/rating column.
    if "%" in s or "industry" in s or "rating" in s:
        return False
    return "instrument" in s or "issuer" in s


def _is_industry_header(s: str) -> bool:
    return ("industry" in s or "rating" in s) and "%" not in s


def classify_section(label: str) -> str | None:
    """Map a section-banner string to an instrument_type, or None if the
    label is not a recognized section banner (caller keeps prior section).

    Shared SEBI-portfolio banner vocabulary across AMCs. Subtotal / total /
    grand-total footers return None so they don't reset the section.
    """
    if not label:
        return None
    s = label.strip().lower()
    if s in ("subtotal", "sub total", "total", "grand total", "sub-total"):
        return None
    if "equity" in s and "derivative" not in s:
        return "Equity"
    if any(k in s for k in (
        "debt", "money market", "triparty repo", "tri-party repo", "treps",
        "treasury", "government securities", "g-sec", "bonds", "debentures",
        "certificate of deposit", "commercial paper", "zero coupon",
        "preference shares", "non convertible", "securitised debt",
        "corporate debt", "fixed deposit",
    )):
        return "Debt"
    if "reit" in s or "invit" in s:
        return "REIT/InvIT"
    if any(k in s for k in (
        "net current asset", "net receivable", "net payable", "cash",
        "tbill", "t-bill", "cash margin", "cash & cash equivalent",
    )):
        return "Cash"
    return None


# --- Statement-date ("AS ON <date>") banner detection -----------------------
#
# SEBI portfolio workbooks print the statement date in a pre-header banner
# row ("MONTHLY PORTFOLIO STATEMENT AS ON 30 Apr 2026"). We parse it to
# validate the artifact's INTERNAL month against the month it was requested
# as — an endpoint serving last month's file for this month's id (quant,
# 2026-05) must be rejected, not cached.

_MONTH_NUM = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5,
    "june": 6, "july": 7, "august": 8, "september": 9, "october": 10,
    "november": 11, "december": 12,
}

_AS_ON_RE = re.compile(r"\bas\s+on\b", re.IGNORECASE)

# Date snippet following an 'as on' phrase. Three shapes:
#   day-first:   '30 Apr 2026', '30-Apr-2026', '30-APR-2026', '30th April 2026'
#   month-first: 'April 30, 2026', 'April 30,2026'
#   numeric day-first (tata, B2 calibration over the live cache): 'as on
#     30-04-2026', 'Portfolio as on 31/05/26', '31-05-26'. Indian convention
#     is strictly dd-mm; statement dates are month-ends (28-31) so the
#     day-first read is never ambiguous in practice. Month is validated
#     1..12 at parse time; the lookarounds stop us matching INSIDE a longer
#     digit run (e.g. the '26-04-30' tail of an ISO '2026-04-30').
# Separators are spaces / hyphens / commas (word-month) or - / . (numeric);
# the day may carry an ordinal suffix; trailing footnote markers
# ('29 May 2026*') are simply not consumed.
_STMT_DATE_RE = re.compile(
    r"""
    (?:
        (?P<d1>\d{1,2})(?:st|nd|rd|th)?[\s\-,]+(?P<m1>[A-Za-z]{3,9})[\s\-,]+(?P<y1>\d{4})
      |
        (?P<m2>[A-Za-z]{3,9})[\s\-,]+(?P<d2>\d{1,2})(?:st|nd|rd|th)?[\s\-,]*(?P<y2>\d{4})
      |
        (?<!\d)(?P<d3>\d{1,2})[./-](?P<m3>\d{1,2})[./-](?P<y3>\d{4}|\d{2})(?!\d)
    )
    """,
    re.IGNORECASE | re.VERBOSE,
)


def _month_token_num(token: str) -> int | None:
    """'Apr' / 'APRIL' / 'Sept' → 4 / 4 / 9; non-month words → None."""
    t = token.lower()
    for name, num in _MONTH_NUM.items():
        if name.startswith(t):
            return num
    return None


def _ym_from_text(text: str) -> str | None:
    """First parseable 'AS ON'-style date in ``text`` as 'YYYY-MM', or None."""
    for m in _STMT_DATE_RE.finditer(text):
        if m.group("m3") is not None:  # numeric dd-mm-yy(yy) form
            month = int(m.group("m3"))
            if not 1 <= month <= 12:
                continue  # '13' etc. — digit triple that isn't a date
            year = int(m.group("y3"))
            if year < 100:
                year += 2000  # '26' → 2026 (no pre-2000 statements exist)
            return f"{year:04d}-{month:02d}"
        month = _month_token_num(m.group("m1") or m.group("m2"))
        if month is None:
            continue  # regex shape matched but the word isn't a month
        return f"{int(m.group('y1') or m.group('y2')):04d}-{month:02d}"
    return None


def statement_months_in_text(text: str, *, window: int = 80) -> set[str]:
    """All months ('YYYY-MM') named by 'as on <date>' phrases in free text.

    Each 'as on' occurrence is paired with the first parseable date in the
    following ``window`` characters (close-binding: a faraway date elsewhere
    in the text must not be attributed to this phrase). Shared by the
    in-cell Excel banner scan below and the managers/factsheet advisory
    first-page scan (B2).
    """
    found: set[str] = set()
    for m in _AS_ON_RE.finditer(text):
        ym = _ym_from_text(text[m.end(): m.end() + window])
        if ym is not None:
            found.add(ym)
    return found


def _ym_from_cell(cell: object) -> str | None:
    """'YYYY-MM' from a date-bearing cell: a datetime/date value (openpyxl
    parses typed date cells) or a string carrying a parseable date."""
    if isinstance(cell, _dt.date):  # covers datetime too
        return f"{cell.year:04d}-{cell.month:02d}"
    if isinstance(cell, str):
        return _ym_from_text(cell)
    return None


def find_statement_months(rows: list[tuple]) -> set[str]:
    """Scan banner rows for 'as on <date>' phrases; return months as 'YYYY-MM'.

    Callers pass the PRE-HEADER rows only (``rows[:header_idx]``) so table
    data is never scanned. Handles the in-cell formats verified across the
    live cache ('30 Apr 2026', 'April 30, 2026', 'April 30,2026',
    '30-Apr-2026', '30th April 2026', '29 May 2026*') and the split-cell
    variant where the cell is just 'AS ON :' and the date — a string or a
    typed datetime cell — sits in a later cell of the same row. An 'as on'
    with no parseable date nearby (e.g. 'NAV As on Record Date') yields
    nothing.
    """
    found: set[str] = set()
    for row in rows:
        for j, cell in enumerate(row):
            if not isinstance(cell, str) or not _AS_ON_RE.search(cell):
                continue
            yms = statement_months_in_text(cell)
            if not yms:
                # Split-cell: date lives in a later cell of the same row.
                for later in row[j + 1:]:
                    ym = _ym_from_cell(later)
                    if ym is not None:
                        yms = {ym}
                        break
            found |= yms
    return found


# Magic bytes distinguishing the two workbook containers we can scan.
_XLSX_MAGIC = b"PK\x03\x04"  # OOXML zip (.xlsx)
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0"  # legacy compound file (.xls)


def _xlsx_head_rows(path: Path, max_rows: int) -> list[list[tuple]]:
    """First ``max_rows`` rows of every sheet of an .xlsx, one list per sheet."""
    out: list[list[tuple]] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            out.append(list(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True)))
    finally:
        wb.close()
    return out


def _xls_head_rows(path: Path, max_rows: int) -> list[list[tuple]]:
    """First ``max_rows`` rows of every sheet of a legacy OLE2 .xls (xlrd).

    Typed date cells are converted to datetimes so the split-cell banner
    variant ('AS ON :' + a date cell) parses the same as on the openpyxl
    path.
    """
    import xlrd  # localized: only OLE2 artifacts need it

    out: list[list[tuple]] = []
    wb = xlrd.open_workbook(str(path))
    for ws in wb.sheets():
        rows: list[tuple] = []
        for i in range(min(max_rows, ws.nrows)):
            cells: list[object] = []
            for cell in ws.row(i):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        cells.append(
                            xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode)
                        )
                    except Exception:  # noqa: BLE001 — malformed date serial
                        cells.append(None)
                else:
                    cells.append(cell.value)
            rows.append(tuple(cells))
        out.append(rows)
    return out


def artifact_statement_months(path: Path, max_rows: int = 15) -> set[str] | None:
    """Union of 'AS ON <date>' banner months over the first ``max_rows`` rows
    of EVERY sheet in the workbook at ``path``.

    This is the orchestrator's central statement-date screen (B2): it
    format-sniffs the artifact (xlsx zip via openpyxl, legacy OLE2 .xls via
    xlrd) so every holdings adapter — including the bespoke fixed-offset
    parsers that never call ``parse_sebi_excel`` — flows through the same
    wrong-month check.

    Returns ``None`` when the artifact cannot be scanned (unknown magic,
    corrupt/non-workbook container such as uti's .zip): 'cannot check' must
    be distinguished from 'checked, found no banner' (empty set). Scanning
    is validate-when-present — the CALLER decides what an empty set means.
    """
    try:
        with open(path, "rb") as fh:
            magic = fh.read(4)
    except OSError:
        return None
    try:
        if magic == _XLSX_MAGIC:
            sheets = _xlsx_head_rows(path, max_rows)
        elif magic == _OLE2_MAGIC:
            sheets = _xls_head_rows(path, max_rows)
        else:
            return None
    except Exception:  # noqa: BLE001 — corrupt workbook: cannot check
        return None
    found: set[str] = set()
    for rows in sheets:
        found |= find_statement_months(rows)
    return found


def _detect_header(rows: list[tuple], max_scan: int = 40) -> tuple[int, dict] | None:
    """Find the header row + column map. Returns (header_row_idx, colmap) or
    None. colmap keys: isin, name, weight, industry (industry optional)."""
    for i, row in enumerate(rows[:max_scan]):
        isin_col = name_col = weight_col = industry_col = None
        for j, cell in enumerate(row):
            s = _norm(cell)
            if not s:
                continue
            if isin_col is None and _is_isin_header(s):
                isin_col = j
            elif weight_col is None and _is_weight_header(s):
                weight_col = j
            elif name_col is None and _is_name_header(s):
                name_col = j
            elif industry_col is None and _is_industry_header(s):
                industry_col = j
        # A valid SEBI portfolio header has at minimum ISIN + weight + name.
        if isin_col is not None and weight_col is not None and name_col is not None:
            return i, {
                "isin": isin_col,
                "name": name_col,
                "weight": weight_col,
                "industry": industry_col,
            }
    return None


def _parse_one_sheet(
    ws,
    *,
    expect_ym: str | None = None,
    require_statement_date: bool = False,
    artifact_path: Path | None = None,
) -> list[ParsedHoldingRecord] | None:
    """Parse a single worksheet in the generic SEBI layout. Returns the
    holding records (unscaled-then-scaled) or None if no header was found
    (i.e. this sheet is not a portfolio table).

    When ``expect_ym`` ('YYYY-MM') is given, the pre-header banner rows are
    scanned for 'AS ON <date>' statement dates; any banner month matching
    ``expect_ym`` passes (tolerates e.g. a lagging riskometer date alongside
    the true portfolio banner), but banners that ALL disagree raise
    ``StatementDateMismatchError`` — a wrong-month artifact must never be
    parsed into the requested month. With ``require_statement_date`` a
    missing banner also raises (for AMCs whose banner is verified always
    present, absence itself is suspicious).
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None
    detected = _detect_header(rows)
    if detected is None:
        return None
    header_idx, cm = detected
    if expect_ym:
        found = find_statement_months(rows[:header_idx])
        if (found and expect_ym not in found) or (
            not found and require_statement_date
        ):
            raise StatementDateMismatchError(artifact_path, expect_ym, found)
    isin_c, name_c, weight_c = cm["isin"], cm["name"], cm["weight"]

    current_section = "Equity"
    raw: list[dict] = []
    for row in rows[header_idx + 1:]:
        ncols = max(isin_c, name_c, weight_c) + 1
        cells = list(row) + [None] * max(0, ncols - len(row))
        isin = cells[isin_c]
        name = cells[name_c]
        weight = cells[weight_c]
        isin_s = str(isin).strip() if isin is not None else ""
        name_s = name.strip() if isinstance(name, str) else ""

        if name_s.lower() == "grand total":
            break
        # Section banner: name present but ISIN col is not an ISIN.
        if name_s and not is_isin(isin_s):
            section = classify_section(name_s)
            if section:
                current_section = section
            continue
        if not is_isin(isin_s) or not name_s:
            continue
        try:
            w = float(weight)
        except (TypeError, ValueError):
            continue
        raw.append({
            "name": name_s.rstrip(" *"),
            "isin": isin_s,
            "weight": w,
            "section": current_section,
        })

    if not raw:
        return None

    # Weight-unit detection: a fraction-format portfolio's weights sum to
    # ~1.0 across ALL instrument types and never exceed ~1.0 in total; a
    # percent-format portfolio sums to ~100. So total ISIN-weight magnitude
    # > 1.5 ⇒ percent. This correctly classifies cash-heavy debt funds whose
    # few ISIN rows still sum > 1.5 in percent (a max-based or sum<5 rule
    # would misread those as fractions). The only residual ambiguity — a
    # fund with <1.5% total in ISIN-bearing instruments — is never an equity
    # fund, so it's irrelevant to the ranked universe.
    total = sum(abs(r["weight"]) for r in raw)
    scale = 100.0 if total <= 1.5 else 1.0

    out: list[ParsedHoldingRecord] = []
    seen: set[str] = set()
    for r in raw:
        if r["isin"] in seen:
            continue
        seen.add(r["isin"])
        out.append(ParsedHoldingRecord(
            scheme_name_printed="",  # filled by caller
            security_name=r["name"],
            weight_pct=r["weight"] * scale,
            isin=r["isin"],
            instrument_type=r["section"],
            source_amc="",  # filled by caller
        ))
    return out


def parse_sebi_excel(
    excel_path: Path,
    scheme_name_printed: str,
    source_amc: str,
    sheet_index: int | None = 0,
    *,
    expect_ym: str | None = None,
    require_statement_date: bool = False,
) -> Iterable[ParsedHoldingRecord]:
    """Parse a single-scheme SEBI monthly-portfolio Excel generically.

    Detects the header row, ISIN/name/weight columns, and the weight unit.
    Yields ISIN-bearing ``ParsedHoldingRecord`` rows with ``security_name``,
    ``weight_pct`` (always percent), ``isin``, and ``instrument_type``.

    ``sheet_index``: which sheet holds the portfolio. Default 0 (the common
    case). Pass ``None`` to scan ALL sheets and parse the first one that
    yields a recognizable SEBI header (useful when the equity table isn't on
    sheet 0). Multi-scheme consolidated workbooks need a bespoke adapter —
    this helper assumes one scheme per file.

    ``expect_ym`` ('YYYY-MM'): validate the workbook's printed 'AS ON' banner
    month against the requested data month; raise
    ``StatementDateMismatchError`` on disagreement (validate-when-present —
    a workbook with no banner passes unless ``require_statement_date``).
    """
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_index is None:
            sheets = list(wb.worksheets)
        else:
            sheets = [wb.worksheets[sheet_index]] if sheet_index < len(wb.worksheets) else []
        for ws in sheets:
            recs = _parse_one_sheet(
                ws,
                expect_ym=expect_ym,
                require_statement_date=require_statement_date,
                artifact_path=excel_path,
            )
            if recs:
                for r in recs:
                    yield ParsedHoldingRecord(
                        scheme_name_printed=scheme_name_printed,
                        security_name=r.security_name,
                        weight_pct=r.weight_pct,
                        isin=r.isin,
                        instrument_type=r.instrument_type,
                        source_amc=source_amc,
                    )
                return  # first portfolio sheet wins
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Config-driven adapter base for the common case (one Excel per scheme,
# discoverable from a static disclosures HTML page).
# ---------------------------------------------------------------------------


def discover_xlsx_links(
    page_url: str,
    url_pattern: re.Pattern,
    scheme_from_filename: Callable[[str], str | None],
    *,
    base_url: str | None = None,
) -> dict[str, str]:
    """Scrape a disclosures HTML page for portfolio Excel links.

    Handles the dominant discovery shape: a (possibly SPA-rendered) HTML
    page whose markup embeds absolute or root-relative ``.xls``/``.xlsx``
    URLs. ``url_pattern`` must capture the full URL in group 1 (or group
    'url'); root-relative matches are absolutized against ``base_url`` (or
    the page origin). ``scheme_from_filename`` turns the URL's filename into
    the printed scheme name (return None to skip a file, e.g. aggregate
    "all-schemes" workbooks). First occurrence of a scheme name wins.
    """
    html = fetch_bytes(page_url).decode("utf-8", errors="replace")
    if base_url is None:
        m = re.match(r"(https?://[^/]+)", page_url)
        base_url = m.group(1) if m else ""
    has_named = "url" in url_pattern.groupindex
    out: dict[str, str] = {}
    for mt in url_pattern.finditer(html):
        url = mt.group("url") if has_named else mt.group(1)
        if url.startswith("/"):
            url = urljoin(base_url + "/", url.lstrip("/"))
        filename = unquote(url.rsplit("/", 1)[-1].split("?", 1)[0])
        scheme = scheme_from_filename(filename)
        if not scheme:
            continue
        out.setdefault(re.sub(r"\s+", " ", scheme).strip(), url)
    return out


class GenericHoldingsAdapter(HoldingsAdapter):
    """Base for AMCs that publish one SEBI-format Excel per scheme.

    Subclasses set ``amc_slug`` / ``source_label`` and implement
    ``discover_scheme_urls`` (often a one-liner via ``discover_xlsx_links``).
    ``parse_excel`` and ``fetch_excel`` are inherited — the shared
    ``parse_sebi_excel`` auto-detects columns + weight unit, and fetch uses
    the standard cached download. Override ``sheet_index = None`` when the
    portfolio table is not on the first sheet. Set
    ``require_statement_date = True`` when the AMC's 'AS ON' banner is
    verified always present, so a banner-less workbook is rejected rather
    than trusted.
    """

    sheet_index: int | None = 0
    require_statement_date: bool = False

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Download one scheme's portfolio Excel; return the local Path.

        Re-uses an existing cached file by mere existence; a forced run
        (``--full`` pipeline / ``--force`` CLI) deletes the data month's
        cached file in ``holdings._run.run_for_amc`` BEFORE calling this,
        so force re-downloads (B9).
        """
        out = paths.holdings_excel_raw(self.amc_slug, ym, scheme_filename)
        if out.exists():
            return out
        # expect='excel' validates magic bytes (xlsx zip OR legacy OLE2 .xls)
        # before caching; an HTML error/WAF page served as 200 raises
        # IngestError and is never written.
        return download_to(url, out, expect="excel")

    def parse_excel(
        self, excel_path: Path, scheme_name_printed: str, ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        return parse_sebi_excel(
            excel_path, scheme_name_printed, self.amc_slug,
            sheet_index=self.sheet_index,
            expect_ym=ym,
            require_statement_date=self.require_statement_date,
        )
