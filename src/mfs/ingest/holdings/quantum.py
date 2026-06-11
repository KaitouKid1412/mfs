"""Quantum Mutual Fund monthly portfolio holdings adapter (Phase 5).

Quantum publishes ONE consolidated workbook per month covering every scheme,
exactly like Nippon. The "combined portfolio" disclosures page is a server-
rendered page (NOT a SPA — the .xlsx links are already in the static HTML)
whose download anchors point at opaque GUID URLs on the AMC's file CDN:

    https://www.quantumamc.com/FileCDN/FactSheet/<guid>.xlsx

The GUID encodes nothing about the data month, so we cannot pattern the URL
by ``ym`` directly. Instead, each anchor carries a Google-Tag-Manager
``onclick="GTMcodeforxml(<url>, <page>, '<Month> <Year> - All Funds', ...)"``
in which the human-readable ``'<Month> <Year>'`` label IS the data month
(the portfolio "as on" month-end). We extract (label, url) pairs from those
GTM calls and pick the one whose label matches the requested ``ym``. This is
robust month-over-month: next month's file simply appears as a new anchor
with a new GUID and a new label, and the same extraction finds it.

Disclosure page (combined / all-funds, all months):
    https://www.quantumamc.com/portfolio/combined/-1/1/0/0

Workbook structure (validated against the April-2026 file):
- Sheet 0 ``Index``: rows 10+ map col B = "Scheme Full Name" to col C =
  "Scheme Code" (e.g. "Quantum Value Fund" -> "QLTEVF"). Header row 9
  reads "Scheme Full Name | Scheme Code". (Note the column order is the
  inverse of Nippon's Index, where the code came first.)
- Sheets 1..N: one sheet per scheme, sheet name == the scheme code.
- Per-scheme layout (validated against QLTEVF / Quantum Value Fund):
    row 10: header — col B="Name of Instrument", col C="ISIN",
            col D="Industry +", col E="Quantity",
            col F="Market/ Fair Value ( Rs. in Lakhs)", col G="% to NAV".
    section banners ("EQUITY & EQUITY RELATED", "Listed /Awaiting listing
    on Stock Exchanges", "MONEY MARKET INSTRUMENTS", "TREPS ^",
    "OTHERS", ...) sit in col B with the ISIN column empty.
    holding rows carry a real ISIN in col C and "% to NAV" in col G.
    A "Grand Total" banner closes the table; a free-text "Notes:" block
    follows it.

CRITICAL UNIT NOTE: like Nippon, Quantum stores "% to NAV" as a FRACTION
(HDFC Bank = 0.0745 for a 7.45% weight), not a percentage. We don't
hand-handle this — the shared ``parse_sebi_excel`` detects the
fraction-vs-percent unit from the per-scheme weight total and scales
fractions up by 100. We locate the scheme's sheet via the Index and
delegate per-sheet parsing to ``parse_sebi_excel(..., sheet_index=...)``,
which also auto-detects the ISIN / name / weight columns and stops at
"Grand Total".
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import openpyxl

from mfs import paths
from mfs.ingest.holdings._generic import GenericHoldingsAdapter, parse_sebi_excel
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to, fetch_bytes
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_DISCLOSURE_PAGE = "https://www.quantumamc.com/portfolio/combined/-1/1/0/0"

_MONTH_NAMES = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)
_MONTH_TO_INT = {name.lower(): i + 1 for i, name in enumerate(_MONTH_NAMES)}

# Each download anchor carries a GTM onclick of the form:
#   onclick="GTMcodeforxml('<xlsx-url>', '<page-url>',
#                          '<Month> <Year> - All Funds', 'All Funds Portfolio');"
# The third argument's '<Month> <Year>' is the data month (portfolio "as on"
# month-end). We capture the xlsx URL (group 'url') and that label (group
# 'label'). Keeping the match anchored to the GTM call avoids associating the
# wrong month with a URL when several .xlsx links sit adjacent in the markup.
_GTM_RE = re.compile(
    r"GTMcodeforxml\(\s*"
    r"'(?P<url>https://www\.quantumamc\.com/FileCDN/FactSheet/[^']+?\.xlsx)'\s*,\s*"
    r"'[^']*'\s*,\s*"
    r"'(?P<label>[A-Za-z]+ \d{4}) - All Funds'",
    re.IGNORECASE,
)


def _label_to_ym(label: str) -> str | None:
    """'April 2026' -> '2026-04'. Returns None for an unparseable label."""
    parts = label.strip().split()
    if len(parts) != 2:
        return None
    month = _MONTH_TO_INT.get(parts[0].lower())
    if month is None:
        return None
    try:
        year = int(parts[1])
    except ValueError:
        return None
    return f"{year:04d}-{month:02d}"


def _master_excel_filename(ym: str) -> str:
    """Canonical local-cache filename for the consolidated monthly workbook."""
    return f"Quantum-MONTHLY-PORTFOLIO-{ym}.xlsx"


def _extract_scheme_name(raw: object) -> str | None:
    """Normalize an Index-sheet 'Scheme Full Name' cell to the printed name."""
    if not isinstance(raw, str):
        return None
    name = re.sub(r"\s+", " ", raw).strip()
    return name or None


@register_adapter
class QuantumHoldingsAdapter(GenericHoldingsAdapter):
    amc_slug = "quantum"
    source_label = "Quantum Mutual Fund"

    # ------------------------------------------------------------------
    # Discovery
    # ------------------------------------------------------------------

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Find the consolidated workbook for ``ym``, cache it once, and
        enumerate the per-scheme entries from its Index sheet.

        Because Quantum publishes a SINGLE workbook for all schemes, every
        returned entry maps to the same URL. ``fetch_excel`` (overridden
        below) short-circuits to the already-cached master file regardless
        of the per-scheme filename it is handed.
        """
        html = fetch_bytes(_DISCLOSURE_PAGE).decode("utf-8", errors="replace")
        master_url = self._select_master_url(html, ym)
        if master_url is None:
            log.warning("holdings.quantum.no_url_for_ym", ym=ym)
            return {}

        master_path = self._cached_master_path(ym)
        if not master_path.exists():
            download_to(master_url, master_path)

        try:
            wb = openpyxl.load_workbook(
                master_path, read_only=True, data_only=True
            )
        except Exception as e:  # noqa: BLE001 — surface any payload issue
            log.error(
                "holdings.quantum.workbook_open_failed",
                ym=ym, path=str(master_path), err=str(e),
            )
            return {}

        out: dict[str, str] = {}
        try:
            if "Index" not in wb.sheetnames:
                log.warning(
                    "holdings.quantum.no_index_sheet",
                    ym=ym, sheetnames=wb.sheetnames[:5],
                )
                return {}
            idx_ws = wb["Index"]
            sheet_set = set(wb.sheetnames)
            for name, code in self._index_rows(idx_ws):
                if code not in sheet_set:
                    # Index points at a sheet that isn't in the workbook —
                    # Quantum-side hiccup; skip rather than fabricate.
                    continue
                out.setdefault(name, master_url)
        finally:
            wb.close()

        log.info(
            "holdings.quantum.discover",
            ym=ym, n_schemes=len(out), master_url=master_url,
        )
        return out

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _index_rows(idx_ws) -> Iterable[tuple[str, str]]:
        """Yield (scheme_full_name, scheme_code) from the Index sheet.

        The Index header row reads "... | Scheme Full Name | Scheme Code"
        with the name in col B (index 1) and the code in col C (index 2).
        We locate the header by name so a layout shift in the preamble
        rows (banner / address block) doesn't throw the offsets off.
        """
        name_col = code_col = None
        for row in idx_ws.iter_rows(values_only=True):
            cells = list(row)
            if name_col is None:
                for j, cell in enumerate(cells):
                    if not isinstance(cell, str):
                        continue
                    s = cell.strip().lower()
                    if s in ("scheme full name", "scheme name"):
                        name_col = j
                    elif s == "scheme code":
                        code_col = j
                if name_col is not None and code_col is not None:
                    continue  # header found; subsequent rows are data
                continue
            if code_col is None:
                continue
            name = _extract_scheme_name(
                cells[name_col] if name_col < len(cells) else None
            )
            code = cells[code_col] if code_col < len(cells) else None
            if not name or code is None:
                continue
            code_s = str(code).strip()
            if not code_s:
                continue
            yield name, code_s

    def _select_master_url(self, html: str, ym: str) -> str | None:
        """Pick the workbook URL whose GTM month-label matches data month ``ym``."""
        for m in _GTM_RE.finditer(html):
            if _label_to_ym(m.group("label")) == ym:
                return m.group("url")
        return None

    def _cached_master_path(self, ym: str) -> Path:
        """Path under data/raw/holdings/quantum/<ym>/ for the master workbook."""
        return paths.holdings_excel_raw(
            self.amc_slug, ym, _master_excel_filename(ym)
        )

    # ------------------------------------------------------------------
    # Download
    # ------------------------------------------------------------------

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        """Return the cached consolidated workbook path for this month.

        Quantum's master Excel is shared across all schemes in ``ym``, so
        ``scheme_filename`` is intentionally ignored. We still download here
        (idempotently) in case the caller invoked fetch without going through
        discover first (e.g. a test).
        """
        master_path = self._cached_master_path(ym)
        if master_path.exists():
            return master_path
        return download_to(url, master_path)

    # ------------------------------------------------------------------
    # Parse
    # ------------------------------------------------------------------

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        """Locate the scheme's sheet inside the consolidated workbook and
        delegate per-sheet parsing to the shared ``parse_sebi_excel``.

        ``parse_sebi_excel`` auto-detects the header row, the ISIN / name /
        weight columns, and the fraction-vs-percent unit, and stops at the
        "Grand Total" banner — so all we owe it is the right ``sheet_index``.
        """
        sheet_index = self._sheet_index_for_scheme(
            excel_path, scheme_name_printed
        )
        if sheet_index is None:
            log.warning(
                "holdings.quantum.scheme_not_in_index",
                scheme=scheme_name_printed, ym=ym,
            )
            return iter(())
        return parse_sebi_excel(
            excel_path, scheme_name_printed, self.amc_slug,
            sheet_index=sheet_index,
        )

    def _sheet_index_for_scheme(
        self, excel_path: Path, scheme_name_printed: str
    ) -> int | None:
        """Resolve a printed scheme name to its 0-based sheet index via the
        Index sheet's name->code map. Returns None if not found."""
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            if "Index" not in wb.sheetnames:
                return None
            target = re.sub(r"\s+", " ", scheme_name_printed.strip()).lower()
            sheetnames = wb.sheetnames
            for name, code in self._index_rows(wb["Index"]):
                if name.lower() == target and code in sheetnames:
                    return sheetnames.index(code)
        finally:
            wb.close()
        return None
