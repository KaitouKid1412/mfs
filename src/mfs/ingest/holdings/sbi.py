"""SBI Mutual Fund monthly portfolio holdings adapter (Phase 3.C).

SBI publishes one Excel per scheme per month. Unlike HDFC (where the page
HTML directly carries every .xlsx URL), SBI's portfolios page is a
Sitefinity SPA that loads the table via an AJAX POST. Discovery uses that
AJAX endpoint:

    POST https://www.sbimf.com/ajaxcall/CMS/GetSchemePortfolioSheets
    body: {"FundId":0,"PSYear":"<YYYY>","PSMonth":"<MonthName>","PSFrequency":"Monthly"}

The response is an HTML <tbody> fragment with one row per scheme; each row
has two <a> tags pointing at the same Excel, with the first carrying the
scheme name as link text.

URL pattern (verified for April 2026):
    https://www.sbimf.com/docs/default-source/scheme-portfolios/
      <scheme-slug>-monthly-portfolio---<month-lower>-<YYYY>.xlsx[?sfvrsn=...]

Excel layout (validated against SBI Large Cap Fund, April 2026):
- Single sheet named after the scheme code (e.g. 'SBLUECHIP').
- Row 1-5: scheme banner + "PORTFOLIO STATEMENT AS ON :" timestamp.
- Row 6 is the column header:
    [B='' | C='Name of the Instrument / Issuer' | D='ISIN' |
     E='Rating / Industry^' | F='Quantity' | G='Market value (Rs. in Lakhs)' |
     H='% to AUM' | I='YTM %' | J='YTC %'].
- Row 7 blank, then row 8+ section banners + holdings interleaved:
    'EQUITY & EQUITY RELATED'  (col C)
    'a) Listed/awaiting listing on Stock Exchanges' (col C)
    'Equity Shares'  (col C)
    <holding rows: col B = internal code, C = name, D = ISIN, etc.>
- Banner rows have label in col C and other cells empty/None.

Note: SBI's column positions differ from HDFC. HDFC had ISIN in col B; SBI
has ISIN in col D. We don't try to share `parse_excel` between adapters —
each AMC writes its own column-aware parser.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import openpyxl

from mfs import paths
from mfs.io.http import download_to, fetch_bytes
from mfs.ingest.holdings._base import HoldingsAdapter
from mfs.ingest.holdings._registry import register_adapter
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

_LISTING_API = "https://www.sbimf.com/ajaxcall/CMS/GetSchemePortfolioSheets"
_FILES_HOST = "www.sbimf.com"

_MONTH_FULL = (
    "January|February|March|April|May|June|"
    "July|August|September|October|November|December"
)

# One anchor per scheme. The body fragment has two anchors per row (a label
# and a Download button); we only keep the labeled one so we don't double-count.
_ANCHOR_RE = re.compile(
    r'<a\s+href="(?P<url>https://www\.sbimf\.com/docs/default-source/'
    r'scheme-portfolios/[^"]+?\.xlsx(?:\?[^"]*)?)"\s+target="_blank">'
    r'(?P<label>[^<]+)</a>',
    re.IGNORECASE,
)

# Strict regex for per-scheme Excel filenames. Excludes the aggregate
# "all-schemes-monthly-portfolio" file (the slug doesn't start with `all-`).
_FILENAME_RE = re.compile(
    rf"^(?P<slug>(?:sbi|magnum)-[a-z0-9\-]+?)"
    r"-monthly-portfolio---"
    rf"(?P<month>{_MONTH_FULL.lower()})-(?P<year>\d{{4}})\.xlsx$",
    re.IGNORECASE,
)


def _strip_url_query(url: str) -> str:
    """Drop the cache-buster ?sfvrsn=... — the bare URL works."""
    return url.split("?", 1)[0]


def _decode_url_path(raw: str) -> str:
    from urllib.parse import unquote
    return unquote(raw)


def _data_month_name(data_ym: str) -> tuple[str, str]:
    """data_ym='2026-04' → ('April', '2026'). SBI keys the listing API by
    data-month-name + year, not by publish month."""
    y, m = map(int, data_ym.split("-"))
    return _MONTH_FULL.split("|")[m - 1], f"{y:04d}"


def _is_isin(s: str) -> bool:
    return bool(_ISIN_RE.match(s))


_ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}\d$")


def _classify_section(label: str) -> str | None:
    """Map a section banner to an instrument_type. SBI uses slightly
    different banner wording than HDFC: 'EQUITY & EQUITY RELATED' /
    'a) Listed/awaiting listing on Stock Exchanges' / 'Equity Shares' /
    'DEBT INSTRUMENTS' / 'MONEY MARKET INSTRUMENTS' / 'TREPS' / 'REITs &
    InvITs' / 'Net Receivables/(Payables)'."""
    if not label:
        return None
    s = label.strip().lower()
    if "equity" in s and "derivative" not in s:
        return "Equity"
    if any(k in s for k in (
        "debt", "money market", "treps", "tri-party", "treasury",
        "government securities", "bonds", "debentures", "commercial paper",
    )):
        return "Debt"
    if "reit" in s or "invit" in s:
        return "REIT/InvIT"
    if any(k in s for k in (
        "net receivable", "net current asset", "cash", "tbill", "t-bill",
        "net payable",
    )):
        return "Cash"
    return None


@register_adapter
class SbiHoldingsAdapter(HoldingsAdapter):
    amc_slug = "sbi"
    source_label = "SBI Mutual Fund"

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        month_name, year = _data_month_name(ym)
        body = (
            '{"FundId":0,"PSYear":"' + year + '","PSMonth":"' + month_name
            + '","PSFrequency":"Monthly"}'
        ).encode("utf-8")
        # The listing API needs a POST with JSON content-type. Use httpx
        # directly because our fetch_bytes is GET-only.
        import httpx
        with httpx.Client(
            timeout=30.0,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"
                ),
                "Content-Type": "application/json;charset=utf-8",
                "X-Requested-With": "XMLHttpRequest",
                "Accept": "*/*",
            },
            follow_redirects=True,
        ) as c:
            r = c.post(_LISTING_API, content=body)
            r.raise_for_status()
            html = r.text
        import html as _html
        out: dict[str, str] = {}
        for m in _ANCHOR_RE.finditer(html):
            url = _strip_url_query(m.group("url"))
            # Decode HTML entities in the label text (SBI uses &amp;).
            label = _html.unescape(m.group("label"))
            label = re.sub(r"\s+", " ", label).strip()
            filename = _decode_url_path(url.rsplit("/", 1)[-1])
            if not _FILENAME_RE.match(filename):
                # Aggregate "all-schemes" or unexpected pattern — skip.
                continue
            # Strip the trailing "MONTHLY PORTFOLIO - <MONTH> <YEAR>" noise
            # from the label to recover the printed scheme name.
            scheme = re.sub(
                r"\s+monthly portfolio.*$", "", label, flags=re.IGNORECASE,
            ).strip()
            if not scheme:
                continue
            out.setdefault(scheme, url)
        log.info(
            "holdings.sbi.discover",
            ym=ym, month=month_name, year=year, n_schemes=len(out),
        )
        return out

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        out = paths.holdings_excel_raw(self.amc_slug, ym, scheme_filename)
        if out.exists():
            return out
        return download_to(url, out)

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        current_section = "Equity"
        for row in ws.iter_rows(values_only=True):
            cells = list(row) + [None] * max(0, 8 - len(row))
            # SBI column layout (0-indexed):
            #   col B (1): internal scheme code (data rows only)
            #   col C (2): security name OR section banner
            #   col D (3): ISIN
            #   col H (7): % to AUM (weight)
            name_or_banner = cells[2]
            isin = cells[3]
            weight = cells[7]
            isin_str = isin.strip() if isinstance(isin, str) else ""
            # Section banner row: text in col C, but no ISIN (col D empty).
            if (
                isinstance(name_or_banner, str)
                and name_or_banner.strip()
                and not isin_str
            ):
                section = _classify_section(name_or_banner)
                if section:
                    current_section = section
                continue
            if not _is_isin(isin_str):
                continue
            if not isinstance(name_or_banner, str) or not name_or_banner.strip():
                continue
            try:
                w = float(weight)
            except (TypeError, ValueError):
                continue
            yield ParsedHoldingRecord(
                scheme_name_printed=scheme_name_printed,
                security_name=name_or_banner.strip(),
                weight_pct=w,
                isin=isin_str,
                instrument_type=current_section,
                source_amc=self.amc_slug,
            )
