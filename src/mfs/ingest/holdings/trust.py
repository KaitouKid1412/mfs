"""Trust Mutual Fund (TRUSTMF) monthly portfolio holdings adapter.

Like Nippon/Tata/Groww (and unlike HDFC's one-Excel-per-scheme), Trust
publishes ONE consolidated workbook per month covering every TRUSTMF scheme
(11 sheets for April 2026, one per scheme code: TMFFLEXI / TMFMID / TMFMCAP /
TMFSCAP / TMFLIQ / …).

Disclosure source — the un-walled CMS JSON API
---------------------------------------------
``www.trustmf.com`` is a Vite/React SPA, and its WordPress media host
(``/trustmfsys/wp-content/uploads/.../*.pdf``) is fronted by a WAF that
returns a 487-byte React SPA shell to programmatic clients (this is what
blocked the factsheet/PTR build — see ``ingest/managers/trust.py``).

The PORTFOLIO Excels live on a DIFFERENT, un-walled path. The SPA's download
section is served by a Cosmos CMS API discovered from the page bundle
(``/config.json`` -> ``{PROTOCOL}://{VITE_COSMOSAPIURL}`` ->
``https://www.trustmf.com/api/api/``). Every listing is a POST to
``…/Trust/GetData`` with a ``{systemQueryFileName, tagName, …}`` body. The
monthly portfolios are returned by::

    POST https://www.trustmf.com/api/api/Trust/GetData
    {"systemQueryFileName":"disclosuresweb.xml",
     "tagName":"GetDisclosureByType",
     "sortField":"uploaddate","sortDirection":"DESC",
     "replaceField":"_slug_","replaceValue":"portfolio-monthly-disclosure"}

-> resultSetArray = [{"title":"TRUSTMF Monthly Portfolio Report as on
   30.04.2026", "fileurl":"https://trustmf.com/Content/2026/5/Monthly
   Port_20260508133215.xlsx", "uploaddate":"5/8/2026 …", …}, …].

The April-2026 ``fileurl`` is served (HTTP 200, ~5 MB OOXML ``.xlsx``) from
the ``/Content/`` path — NOT the walled ``/trustmfsys/wp-content/uploads/``
path. The bare-host ``trustmf.com`` 307-redirects to ``www.trustmf.com``;
our HTTP client follows redirects. We key discovery off the DATA month-end
date parsed from the ``title`` ("as on DD.MM.YYYY") rather than the folder in
the URL (which is the publish month, data + 1) — exactly the Tata pattern.

Older listings (Jan-2026 and earlier) carry a ``fileurl`` that still points
at the walled ``…/content/<yyyy>/<mm>/…wp-content…`` host; we don't special-
case them because the date filter selects only the requested ``ym``, and the
recent months use the un-walled ``/Content/`` path.

Workbook structure (validated against April-2026 file, ~5 MB, 11 sheets):
- One sheet per scheme; sheet names are opaque codes ('TMFFLEXI', 'TMFMID',
  …). There is NO Index sheet — the printed scheme name lives in row 1
  (0-indexed), col B, e.g. 'TRUSTMF Flexi Cap Fund'. (Row 0 col B carries a
  stale template code that is the SAME on every sheet and is ignored.)
- Per-scheme layout is the standard SEBI table: row 5 header
  ['Name of the Instrument' | 'ISIN' | 'Industry' | 'Quantity' |
  'Market/Fair Value (Rs. in Lakhs)' | '% to Net Assets' | YTM | YTC | Notes],
  then 'Equity & Equity related' / '(a) Listed / awaiting listing …' /
  Debt/MoneyMarket section banners interleaved with ISIN rows. '% to Net
  Assets' is stored as a FRACTION (ICICI Bank at 0.0591 == 5.91%).

Because the per-scheme sheet IS the canonical SEBI layout, the parse delegates
to the validated generic single-sheet parser (``parse_sebi_excel``), which
auto-detects the header row, ISIN/name/weight columns, and the percent-vs-
fraction weight unit, dedupes ISINs, and stops at GRAND TOTAL. The only
bespoke plumbing (modeled on the Groww adapter) is the consolidated-workbook
discovery + per-sheet banner enumeration. Validated on TRUSTMF Flexi Cap (63
equity ISIN rows summing to 92.6%), Mid Cap, Multi Cap, and Small Cap.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path

import httpx
import openpyxl

from mfs import paths
from mfs.ingest.holdings._generic import GenericHoldingsAdapter, parse_sebi_excel
from mfs.ingest.holdings._registry import register_adapter
from mfs.io.http import download_to
from mfs.schemas import ParsedHoldingRecord
from mfs.utils.logging import get_logger

log = get_logger(__name__)

# Cosmos CMS API base (from the SPA's /config.json -> {PROTOCOL}://{VITE_COSMOSAPIURL}).
_GETDATA_API = "https://www.trustmf.com/api/api/Trust/GetData"
_DISCLOSURE_SLUG = "portfolio-monthly-disclosure"

_UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)

# "TRUSTMF Monthly Portfolio Report as on 30.04.2026" -> DD.MM.YYYY data date.
_TITLE_DATE_RE = re.compile(r"as on\s+(\d{1,2})\.(\d{1,2})\.(\d{4})", re.IGNORECASE)


def _title_data_ym(title: str) -> str | None:
    """Reverse-derive the DATA month (YYYY-MM) from a listing ``title``.

    Trust labels each file with the data month-end ("as on 30.04.2026"), so
    the title date is authoritative — independent of the publish-month folder
    in the file URL (the folder is data + 1).
    """
    if not title:
        return None
    m = _TITLE_DATE_RE.search(title)
    if not m:
        return None
    _day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if not 1 <= month <= 12:
        return None
    return f"{year:04d}-{month:02d}"


def _master_excel_filename(ym: str) -> str:
    """Canonical local-cache filename for the consolidated monthly workbook."""
    return f"TRUSTMF-Monthly-Portfolio-{ym}.xlsx"


def _scheme_banner(ws) -> str | None:
    """Printed scheme name = row 1 (0-indexed), first non-empty string cell.

    Row 0 carries a stale template code (identical across sheets); row 1
    carries the real name, e.g. 'TRUSTMF Flexi Cap Fund'.
    """
    rows = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    if not rows:
        return None
    for cell in rows[0]:
        if isinstance(cell, str) and cell.strip():
            return re.sub(r"\s+", " ", cell).strip()
    return None


@register_adapter
class TrustHoldingsAdapter(GenericHoldingsAdapter):
    amc_slug = "trust"
    source_label = "Trust Mutual Fund"

    # ------------------------------------------------------------------
    # Discovery
    # ------------------------------------------------------------------

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        """Find the consolidated monthly workbook for data month ``ym`` via the
        Cosmos CMS API, cache it once, and enumerate the per-scheme banners
        from its sheets.

        Every scheme maps to the SAME master URL (consolidated workbook);
        ``parse_excel`` later resolves the printed name back to its sheet.
        """
        master_url = self._select_master_url(ym)
        if master_url is None:
            log.warning("holdings.trust.no_url_for_ym", ym=ym)
            return {}

        master_path = self._cached_master_path(ym)
        if not master_path.exists():
            download_to(master_url, master_path)

        try:
            wb = openpyxl.load_workbook(
                master_path, read_only=True, data_only=True,
            )
        except Exception as e:  # noqa: BLE001 — surface any payload issue
            log.error(
                "holdings.trust.workbook_open_failed",
                ym=ym, path=str(master_path), err=str(e),
            )
            return {}

        out: dict[str, str] = {}
        try:
            for code in wb.sheetnames:
                name = _scheme_banner(wb[code])
                if not name:
                    continue
                out.setdefault(name, master_url)
        finally:
            wb.close()

        log.info(
            "holdings.trust.discover",
            ym=ym, n_schemes=len(out), master_url=master_url,
        )
        return out

    def _select_master_url(self, ym: str) -> str | None:
        """Query the CMS API for monthly portfolios and return the file URL
        whose title-encoded DATA month equals ``ym`` (or None if not yet
        published)."""
        body = {
            "systemQueryFileName": "disclosuresweb.xml",
            "tagName": "GetDisclosureByType",
            "searchField": "",
            "searchValue": "",
            "sortField": "uploaddate",
            "sortDirection": "DESC",
            "replaceField": "_slug_",
            "replaceValue": _DISCLOSURE_SLUG,
        }
        with httpx.Client(
            timeout=60.0,
            follow_redirects=True,
            headers={
                "User-Agent": _UA,
                "Content-type": "application/json; charset=UTF-8",
                "Accept": "application/json, text/plain, */*",
                "Origin": "https://www.trustmf.com",
                "Referer": "https://www.trustmf.com/",
            },
        ) as c:
            r = c.post(_GETDATA_API, json=body)
            r.raise_for_status()
            items = (r.json() or {}).get("resultSetArray") or []

        for item in items:
            if _title_data_ym(item.get("title") or "") != ym:
                continue
            url = (item.get("fileurl") or "").strip()
            if url:
                return url
        return None

    def _cached_master_path(self, ym: str) -> Path:
        return paths.holdings_excel_raw(
            self.amc_slug, ym, _master_excel_filename(ym),
        )

    # ------------------------------------------------------------------
    # Download — every scheme shares the one consolidated workbook.
    # ------------------------------------------------------------------

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        master_path = self._cached_master_path(ym)
        if master_path.exists():
            return master_path
        return download_to(url, master_path)

    # ------------------------------------------------------------------
    # Parse — locate the scheme's sheet, then reuse the generic SEBI parser.
    # ------------------------------------------------------------------

    def parse_excel(
        self,
        excel_path: Path,
        scheme_name_printed: str,
        ym: str,
    ) -> Iterable[ParsedHoldingRecord]:
        sheet_index = self._sheet_index_for_scheme(excel_path, scheme_name_printed)
        if sheet_index is None:
            log.warning(
                "holdings.trust.scheme_sheet_not_found",
                scheme=scheme_name_printed, ym=ym,
            )
            return iter(())
        return parse_sebi_excel(
            excel_path,
            scheme_name_printed,
            self.amc_slug,
            sheet_index=sheet_index,
            expect_ym=ym,
        )

    @staticmethod
    def _sheet_index_for_scheme(
        excel_path: Path, scheme_name_printed: str,
    ) -> int | None:
        """Map a printed scheme name to its sheet index by matching the
        row-1 col-B banner in each sheet."""
        target = re.sub(r"\s+", " ", scheme_name_printed.strip()).lower()
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            for idx, code in enumerate(wb.sheetnames):
                name = _scheme_banner(wb[code])
                if name and name.lower() == target:
                    return idx
        finally:
            wb.close()
        return None
