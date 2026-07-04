"""Runner-level tests for holdings statement-date validation (Urgent-3).

A wrong-month artifact (quant's WebForms endpoint served the APRIL files for
a MAY request) must abort the AMC via ``StatementDateMismatchError`` BEFORE
its DB transaction, and the cached Excel must be evicted — otherwise
``fetch_excel``'s exists() check pins the wrong file forever and a genuine
May publication can never replace it.

No network, no live DB: a stub adapter writes a synthetic April-banner
workbook into tmp_path and the DB read/write seams are monkeypatched
(pattern: tests/test_adapter_isolation.py).
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import polars as pl
import pytest

from mfs.errors import StatementDateMismatchError
from mfs.ingest.holdings import _run as holdings_run
from mfs.ingest.holdings._generic import GenericHoldingsAdapter

_SLUG = "stubamc"
_SCHEME = "Stub Flexi Cap Fund"

# quant-shaped workbook: 'AS ON 30 Apr 2026' banner above a valid SEBI table.
_APRIL_BANNER_ROWS = [
    [_SCHEME, None, None, None],
    [None, None, None, None],
    [None, None, None, None],
    ["MONTHLY PORTFOLIO STATEMENT AS ON 30 Apr 2026", None, None, None],
    ["ISIN", "Name of the Instrument", "Industry", "% to NAV"],
    ["EQUITY & EQUITY RELATED", None, None, None],
    ["INE040A01034", "HDFC Bank Ltd", "Banks", 9.24],
    [None, "Grand Total", None, 100.0],
]


class _StubAprilAdapter(GenericHoldingsAdapter):
    """Generic adapter whose 'download' yields an April-banner workbook —
    simulating an endpoint that serves last month's file for this month's
    id. ``parse_excel`` is the REAL inherited one (expect_ym=ym)."""

    amc_slug = _SLUG
    source_label = "Stub AMC"

    def __init__(self, cache_dir: Path) -> None:
        self._cache_dir = cache_dir
        self.fetched: list[Path] = []

    def discover_scheme_urls(self, ym: str) -> dict[str, str]:
        return {_SCHEME: "http://example.test/stub.xlsx"}

    def fetch_excel(self, url: str, scheme_filename: str, ym: str) -> Path:
        out = self._cache_dir / self.amc_slug / ym / scheme_filename
        out.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        for row in _APRIL_BANNER_ROWS:
            wb.active.append(row)
        wb.save(out)
        self.fetched.append(out)
        return out


def _fake_scheme_master() -> pl.DataFrame:
    return pl.DataFrame({
        "scheme_code": ["STUB001"],
        "scheme_name": [_SCHEME],
        "amc_code": [_SLUG],
        "plan_type": ["DIRECT"],
        "option_type": ["GROWTH"],
        "is_active": [True],
    })


@pytest.mark.db
def test_wrong_month_artifact_aborts_amc_evicts_cache_never_writes(
    tmp_path, monkeypatch,
):
    stub = _StubAprilAdapter(tmp_path)
    monkeypatch.setattr(holdings_run, "get_adapter", lambda slug: stub)
    monkeypatch.setattr(
        holdings_run.q, "scheme_master", lambda **kw: _fake_scheme_master(),
    )
    monkeypatch.setattr(
        holdings_run.w, "upsert_holdings",
        lambda *a, **kw: pytest.fail(
            "upsert_holdings must NOT be called for a wrong-month artifact"
        ),
    )

    with pytest.raises(StatementDateMismatchError) as ei:
        holdings_run.run_for_amc(_SLUG, ym="2026-05")

    # The error carries the diagnosis: requested May, artifact says April.
    assert ei.value.expected_ym == "2026-05"
    assert ei.value.found_yms == {"2026-04"}

    # The wrong-month file WAS downloaded into the month-keyed cache, then
    # evicted by the runner — it must not survive to poison later runs.
    assert stub.fetched == [tmp_path / _SLUG / "2026-05" / f"{_SCHEME}.xlsx"]
    assert not stub.fetched[0].exists()


def test_run_all_isolates_month_mismatch_to_one_amc(tmp_path, monkeypatch):
    # The aborting AMC is recorded as its own failure; other AMCs proceed
    # (per-AMC isolation — the pipeline stage must not crash).
    monkeypatch.setattr(
        holdings_run, "registered_adapters", lambda: ["ok_amc", _SLUG],
    )

    def fake_run(slug, ym=None, force=False):
        if slug == _SLUG:
            raise StatementDateMismatchError(
                tmp_path / "x.xlsx", "2026-05", {"2026-04"},
            )
        return {"amc_slug": slug, "ym": ym, "rows_written": 7}

    monkeypatch.setattr(holdings_run, "run_for_amc", fake_run)

    res = holdings_run.run_all(ym="2026-05")
    assert res["ok_amc"]["rows_written"] == 7
    assert res[_SLUG]["error_type"] == "StatementDateMismatchError"
    assert res[_SLUG]["rows_written"] == 0


# ---------------------------------------------------------------------------
# B2 coverage extension: BESPOKE adapters (own parse_excel, never calling
# parse_sebi_excel — the hdfc/sbi/nippon/icici_pru shape) must flow through
# the orchestrator's central artifact statement-date screen.
# ---------------------------------------------------------------------------


class _StubBespokeAdapter(_StubAprilAdapter):
    """Bespoke-shaped adapter: parse_excel is hand-rolled (fixed offsets,
    no expect_ym, no parse_sebi_excel) — exactly the adapters Stage 1 left
    uncovered. Only the orchestrator's central screen can reject its
    wrong-month artifact."""

    def parse_excel(self, excel_path, scheme_name_printed, ym):
        from mfs.schemas import ParsedHoldingRecord

        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            for row in wb.active.iter_rows(values_only=True):
                isin = row[0]
                if not (isinstance(isin, str) and isin.startswith("INE")):
                    continue
                yield ParsedHoldingRecord(
                    scheme_name_printed=scheme_name_printed,
                    security_name=row[1],
                    weight_pct=float(row[3]),
                    isin=isin,
                    instrument_type="Equity",
                    source_amc=self.amc_slug,
                )
        finally:
            wb.close()


@pytest.mark.db
def test_bespoke_adapter_flows_through_central_screen(tmp_path, monkeypatch):
    stub = _StubBespokeAdapter(tmp_path)
    monkeypatch.setattr(holdings_run, "get_adapter", lambda slug: stub)
    monkeypatch.setattr(
        holdings_run.q, "scheme_master", lambda **kw: _fake_scheme_master(),
    )
    monkeypatch.setattr(
        holdings_run.w, "upsert_holdings",
        lambda *a, **kw: pytest.fail(
            "upsert_holdings must NOT be called for a wrong-month artifact"
        ),
    )

    with pytest.raises(StatementDateMismatchError) as ei:
        holdings_run.run_for_amc(_SLUG, ym="2026-05")
    assert ei.value.expected_ym == "2026-05"
    assert ei.value.found_yms == {"2026-04"}
    assert not stub.fetched[0].exists()  # evicted, never pinned


@pytest.mark.db
def test_bespoke_adapter_right_month_parses(tmp_path, monkeypatch):
    stub = _StubBespokeAdapter(tmp_path)
    written = []
    monkeypatch.setattr(holdings_run, "get_adapter", lambda slug: stub)
    monkeypatch.setattr(
        holdings_run.q, "scheme_master", lambda **kw: _fake_scheme_master(),
    )
    monkeypatch.setattr(
        holdings_run.w, "upsert_holdings",
        lambda df, conn=None: written.append(df) or len(df),
    )

    class _Cur:
        def fetchone(self):
            return (0,)

    class _Conn:
        def execute(self, sql, params=None):
            return _Cur()

    from contextlib import contextmanager

    @contextmanager
    def fake_connect(autocommit=False):
        yield _Conn()

    import mfs.db.connection as dbconn
    monkeypatch.setattr(dbconn, "connect", fake_connect)

    res = holdings_run.run_for_amc(_SLUG, ym="2026-04")
    assert res["rows_written"] == 1
    assert written and written[0]["isin"].to_list() == ["INE040A01034"]


@pytest.mark.db
def test_per_adapter_override_hook_wins(tmp_path, monkeypatch):
    # An AMC whose banner is nonstandard can supply its own extraction; the
    # orchestrator must consult it INSTEAD of the generic sniffing scan.
    stub = _StubBespokeAdapter(tmp_path)
    stub.artifact_statement_months = lambda path: {"2026-05"}  # banner lies
    written = []
    monkeypatch.setattr(holdings_run, "get_adapter", lambda slug: stub)
    monkeypatch.setattr(
        holdings_run.q, "scheme_master", lambda **kw: _fake_scheme_master(),
    )
    monkeypatch.setattr(
        holdings_run.w, "upsert_holdings",
        lambda df, conn=None: written.append(df) or len(df),
    )

    class _Cur:
        def fetchone(self):
            return (0,)

    class _Conn:
        def execute(self, sql, params=None):
            return _Cur()

    from contextlib import contextmanager

    @contextmanager
    def fake_connect(autocommit=False):
        yield _Conn()

    import mfs.db.connection as dbconn
    monkeypatch.setattr(dbconn, "connect", fake_connect)

    res = holdings_run.run_for_amc(_SLUG, ym="2026-05")
    assert res["rows_written"] == 1  # override said May; generic scan ignored
